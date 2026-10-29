"""Excel tracker: duplicate check and row writing.

Layout (see SPEC.md): Sheet1, table Table1, header on row 1. Columns are found by their header
name (FIELDS, plus the optional ones left empty), never by letter, so a moved column can't
receive the wrong data. A row's values are a dict keyed by those names.
"""
from __future__ import annotations

import json
import os
import re
import shutil
from copy import copy
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import MultiCellRange

SHEET = "Sheet1"
TABLE = "Table1"
DATE_FORMAT = "M/d/yyyy"
OUTCOME = "Applied"
OUTCOMES = ("Applied", "Rejected", "Awaiting reply")  # the Outcome dropdown
STAGES = ("Online Assessment", "Initial Screening", "Interview")  # the Stage dropdown; blank = no news
AWAITING = "Awaiting reply"  # Outcome whenever a Stage is set
FIELDS = ("Company", "Position", "URL", "Posted date/ Source", "Date", "Outcome", "Stage", "Last Updated",
          "Remarks")  # required headers
OPTIONAL = ("Experience", "Key Skills")  # left empty
DATE_FIELDS = ("Date", "Last Updated")
TABLE_GROWTH = 100  # rows added to Table1 and its date checks and dropdowns when the table is full
TRACKING_PARAMS = {"source", "gh_src", "lever-source", "lever-origin",  # plus every utm_*
                   # iCIMS page-display params
                   "mobile", "width", "height", "bga", "needsredirect", "jan1offset", "jun1offset",
                   "in_iframe"}
JOB_PAGE_SUFFIXES = ("/apply", "/application")  # Lever / Ashby form pages = same job


class TrackerLocked(Exception):
    """The tracker file can't be written, almost always because it is open in Excel."""


class TrackerLayoutError(Exception):
    """Row 1 lacks a header the app needs, so nothing is read or written."""


def norm_header(value) -> str:
    return "".join(str(value or "").casefold().split())


def column_map(header) -> dict[str, int]:
    """Header name -> column number (1-based) from row 1's values, ignoring case and spaces.
    Raises TrackerLayoutError naming every missing required header."""
    found = {}
    for col, value in enumerate(header, start=1):
        found.setdefault(norm_header(value), col)
    cols = {f: found[norm_header(f)] for f in FIELDS + OPTIONAL if norm_header(f) in found}
    missing = [f for f in FIELDS if f not in cols]
    if missing:
        raise TrackerLayoutError(f"row 1 is missing the header{'s' if len(missing) > 1 else ''} {', '.join(missing)}")
    return cols


@dataclass
class TrackerRow:
    company: str
    position: str
    url: str
    posted: str  # Posted date/ Source text, e.g. "Posted today-Tsenta"
    applied_on: date

    def to_json(self) -> dict:
        d = asdict(self)
        d["applied_on"] = self.applied_on.isoformat()
        return d

    def values(self) -> dict:
        """The row as it is written, like a row from read_rows()."""
        return {"Company": self.company, "Position": self.position, "URL": self.url,
                "Posted date/ Source": self.posted, "Date": self.applied_on, "Outcome": OUTCOME, "Stage": "",
                "Last Updated": self.applied_on, "Remarks": ""}

    @classmethod
    def from_json(cls, d: dict) -> "TrackerRow":
        return cls(**{**d, "applied_on": date.fromisoformat(d["applied_on"])})


# ---------- duplicate check ----------

def normalize_url(url: str) -> str:
    """Comparable form of a job URL: no scheme, www, fragment, trailing slash or tracking params."""
    parts = urlsplit(url.strip())
    host = (parts.hostname or "").lower().removeprefix("www.")
    path = parts.path.rstrip("/")
    for suffix in JOB_PAGE_SUFFIXES:
        if path.lower().endswith(suffix):
            path = path[: -len(suffix)]
    query = sorted(
        (k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True)
        if not (k.lower().startswith("utm_") or k.lower() in TRACKING_PARAMS)
    )
    return f"{host}{path}" + (f"?{urlencode(query)}" if query else "")


def _norm_text(value) -> str:
    return " ".join(str(value).casefold().split()) if value else ""


def match_rows(rows, url: str | None = None, company: str | None = None,
               position: str | None = None) -> list[tuple[int, dict]]:
    """Pick duplicates from `rows` = [(row number, values, urls)]. -> [(row, values)]

    A row matches if one of its URLs matches `url`, or its company AND position both match.
    """
    want_url = normalize_url(url) if url else None
    want_pair = (_norm_text(company), _norm_text(position)) if company and position else None
    hits = []
    for r, values, urls in rows:
        url_hit = want_url is not None and any(normalize_url(u) == want_url for u in urls)
        pair_hit = want_pair is not None and (_norm_text(values["Company"]),
                                              _norm_text(values["Position"])) == want_pair
        if url_hit or pair_hit:
            hits.append((r, values))
    return hits


def _columns(ws) -> dict[str, int]:
    return column_map(next(ws.iter_rows(max_row=1, values_only=True), ()))


def _sheet_rows(ws, cols: dict[str, int]) -> list[tuple[int, dict, list[str]]]:
    """Job rows of an open worksheet: [(row number, {field: value}, urls in the URL cell)]. Rows with
    no Company and no URL (stray remarks) are skipped."""
    rows = []
    for row in ws.iter_rows(min_row=2, max_col=max(cols.values())):
        values = {f: row[cols[f] - 1].value for f in FIELDS}
        if values["Company"] is None and values["URL"] is None:
            continue
        link = row[cols["URL"] - 1].hyperlink
        urls = [str(values["URL"])] if values["URL"] else []
        if link is not None and link.target:
            urls.append(link.target)
        rows.append((row[0].row, values, urls))
    return rows


def check_layout(path) -> None:
    """Raise TrackerLayoutError if row 1 of the Excel tracker lacks a required header."""
    wb = load_workbook(path, read_only=True)
    try:
        _columns(wb[SHEET])
    finally:
        wb.close()


def read_rows(path) -> list[tuple[int, dict, list[str]]]:
    """Every job row of the Excel tracker: [(row number, {field: value}, urls in the URL cell)]."""
    wb = load_workbook(path)  # not read_only: hyperlink targets are needed too
    try:
        ws = wb[SHEET]
        return _sheet_rows(ws, _columns(ws))
    finally:
        wb.close()


def find_duplicates(path, url: str | None = None, company: str | None = None,
                    position: str | None = None) -> list[tuple[int, dict]]:
    """Rows whose URL matches `url`, or whose company AND position both match. -> [(row, values)]"""
    return match_rows(read_rows(path), url, company, position)


# ---------- search (find.py) ----------

COMPANY_SUFFIXES = {"inc", "llc", "ltd", "corp", "corporation", "co", "plc", "gmbh", "lp", "llp", "ag", "sa"}


def search_key(value) -> str:
    """Comparable text for search: lowercase, punctuation as spaces, trailing company suffixes
    dropped ('Acme, Inc.' -> 'acme', 'S&P Global' -> 's p global'). A name that is only a suffix stays."""
    words = re.sub(r"[^\w]+", " ", str(value or "").casefold()).split()
    while len(words) > 1 and words[-1] in COMPANY_SUFFIXES:
        words.pop()
    return " ".join(words)


def search_rows(rows, text: str) -> list[tuple[int, dict, list[str]]]:
    """Rows from read_rows() where `text` is in the Company, the Position or a URL."""
    key = search_key(text)
    if not key:
        return []
    raw = " ".join(text.casefold().split())
    squeezed = key.replace(" ", "")
    hits = []
    for r, values, urls in rows:
        in_names = any(key in search_key(values[f]) for f in ("Company", "Position"))
        in_urls = any(raw in u.casefold() or squeezed in u.casefold() for u in urls)
        if in_names or in_urls:
            hits.append((r, values, urls))
    return hits


def parse_date(value) -> date | None:
    """A Date cell as a date: an Excel date, or Sheet text like '6/20/2026'. Anything else: None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def format_note(day: date, note: str, existing) -> str:
    """Remarks after adding `note`: 'existing | MM/DD/YYYY note'. No note = Remarks unchanged."""
    existing = "" if existing is None else str(existing).strip()
    note = note.strip()
    if not note:
        return existing
    dated = f"{day:%m/%d/%Y} {note}"
    return f"{existing} | {dated}" if existing else dated


# ---------- Outcome updates (find.py) ----------

@dataclass
class OutcomeUpdate:
    """A new Outcome (and Stage, and optional note) for one job, found again by its details before
    writing: Sheet and Excel row numbers can differ, and either may have changed since the search."""
    company: str
    position: str
    urls: list[str]  # link text and hyperlink target as found
    applied: str  # the Date as an ISO date, or "" when it couldn't be read
    row: int  # where the search found it; tried first
    outcome: str
    note: str
    noted_on: date
    stage: str = ""  # "" = leave Stage as it is

    def to_json(self) -> dict:
        d = asdict(self)
        d["noted_on"] = self.noted_on.isoformat()
        return d

    @classmethod
    def from_json(cls, d: dict) -> "OutcomeUpdate":
        d = {**d, "noted_on": date.fromisoformat(d["noted_on"])}
        if "stage" not in d and d["outcome"] in STAGES:  # saved before Stage had its own column
            d["stage"], d["outcome"] = d["outcome"], AWAITING
        return cls(**d)

    @classmethod
    def for_row(cls, r: int, values: dict, urls: list[str], outcome: str, note: str,
                noted_on: date, stage: str = "") -> "OutcomeUpdate":
        applied = parse_date(values["Date"])
        return cls(company=str(values["Company"] or ""), position=str(values["Position"] or ""),
                   urls=list(urls), applied=applied.isoformat() if applied else "", row=r, outcome=outcome,
                   note=note, noted_on=noted_on, stage=stage)


def _url_ids(urls) -> set[str]:
    """Normalized links that are real URLs (link text like 'click here' can't identify a job)."""
    return {normalize_url(u) for u in urls if urlsplit(u.strip()).hostname}


def locate(rows, update: OutcomeUpdate) -> int | None:
    """The row number of `update`'s job in rows from read_rows(), or None if it isn't there exactly once.

    The hinted row wins if it still holds the job. Otherwise exactly one row must have the same URL
    (or, with no URL, the same company + position); several are narrowed by company + position, then date.
    """
    want_urls = _url_ids(update.urls)
    want_pair = (_norm_text(update.company), _norm_text(update.position))

    def pair(values):
        return (_norm_text(values["Company"]), _norm_text(values["Position"]))

    def same(values, urls):
        return bool(want_urls & _url_ids(urls)) if want_urls else pair(values) == want_pair

    by_row = {r: (values, urls) for r, values, urls in rows}
    if update.row in by_row and same(*by_row[update.row]):
        return update.row
    hits = [(r, values) for r, values, urls in rows if same(values, urls)]
    if len(hits) > 1 and want_urls:
        hits = [(r, values) for r, values in hits if pair(values) == want_pair] or hits
    if len(hits) > 1 and update.applied:
        hits = [(r, values) for r, values in hits
                if (d := parse_date(values["Date"])) and d.isoformat() == update.applied] or hits
    return hits[0][0] if len(hits) == 1 else None


def _ensure_choices(ws, col: int, values) -> None:
    """Make the list validation on column `col` offer every one of `values`. Only adds; never removes."""
    for dv in ws.data_validations.dataValidation:
        formula = dv.formula1 or ""
        covers = any(rng.min_col <= col <= rng.max_col for rng in dv.sqref.ranges)
        if dv.type != "list" or not covers or not (formula.startswith('"') and formula.endswith('"')):
            continue
        choices = [c for c in formula[1:-1].split(",") if c]
        missing = [v for v in values if v not in choices]
        if missing:
            dv.formula1 = '"' + ",".join(choices + missing) + '"'


def _ensure_dropdowns(ws, cols: dict[str, int]) -> None:
    _ensure_choices(ws, cols["Outcome"], OUTCOMES)
    _ensure_choices(ws, cols["Stage"], STAGES)


def update_outcome(path, updates: list[OutcomeUpdate], backup_dir=None) -> tuple[list, list]:
    """Back up, then set Outcome (and Stage when given) and append each note to Remarks.
    -> ([(row, update) written], [not found])

    A job that isn't in the tracker exactly once is not written. Nothing found = no backup, no save."""
    path = Path(path)
    try:
        wb = load_workbook(path)
    except PermissionError as exc:
        raise TrackerLocked(str(exc)) from exc
    ws = wb[SHEET]
    cols = _columns(ws)
    rows = _sheet_rows(ws, cols)
    written, missing = [], []
    for u in updates:
        r = locate(rows, u)
        if r is None:
            missing.append(u)
            continue
        ws.cell(r, cols["Outcome"]).value = u.outcome
        if u.stage:
            ws.cell(r, cols["Stage"]).value = u.stage
        if u.note.strip():
            remarks = ws.cell(r, cols["Remarks"])
            remarks.value = format_note(u.noted_on, u.note, remarks.value)
        written.append((r, u))
    if written:
        _ensure_dropdowns(ws, cols)
        backup(path, backup_dir)
        _save(wb, path)
    wb.close()
    return written, missing


# ---------- writing ----------

def next_row(ws, company_col: int = 1) -> int:
    """First row after the last row that has a Company value."""
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, company_col).value not in (None, ""):
            return r + 1
    return 2


def format_posted(posted_date: date | None, source: str, today: date) -> str:
    """Posted date/ Source: 'Posted today-Tsenta' or 'Posted October 8, 2026-LinkedIn'."""
    if posted_date is None:
        return source
    if posted_date == today:
        return f"Posted today-{source}"
    return f"Posted {posted_date:%B} {posted_date.day}, {posted_date.year}-{source}"


def backup(path, backup_dir=None) -> Path:
    """Copy the tracker to backups/tracker-YYYYMMDD-HHMMSS.xlsx before any write."""
    path = Path(path)
    backup_dir = Path(backup_dir) if backup_dir else path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = backup_dir / f"tracker-{stamp}.xlsx"
    n = 1
    while target.exists():  # two writes in the same second
        target = backup_dir / f"tracker-{stamp}-{n}.xlsx"
        n += 1
    shutil.copy2(path, target)
    return target


def _extend_table(ws, row: int, cols: dict[str, int]) -> None:
    """Grow Table1 and the data validations on the date, Outcome and Stage columns so `row` keeps the
    date checks and the dropdowns."""
    table = ws.tables[TABLE]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    if row <= max_row:
        return
    new_end = row + TABLE_GROWTH
    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_end}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    grown = {cols[f] for f in (*DATE_FIELDS, "Outcome", "Stage")}
    for dv in ws.data_validations.dataValidation:
        ranges = []
        for rng in dv.sqref.ranges:
            if rng.min_col == rng.max_col and rng.min_col in grown and rng.max_row == max_row:
                letter = get_column_letter(rng.min_col)
                ranges.append(f"{letter}{rng.min_row}:{letter}{new_end}")
            else:
                ranges.append(rng.coord)
        dv.sqref = MultiCellRange(" ".join(ranges))


def _save(wb, path: Path) -> None:
    """Save to a temp file next to the tracker, then swap it in, so a failed save can't corrupt it."""
    tmp = path.with_name(f"~{path.stem}.tmp.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:
        raise TrackerLocked(str(exc)) from exc
    finally:
        if tmp.exists():
            tmp.unlink()


def append_row(path, row: TrackerRow, backup_dir=None) -> int:
    """Back up, then write `row` after the last Company row. Returns the row number written."""
    path = Path(path)
    try:
        wb = load_workbook(path)
    except PermissionError as exc:
        raise TrackerLocked(str(exc)) from exc
    ws = wb[SHEET]
    cols = _columns(ws)
    r = next_row(ws, cols["Company"])
    _extend_table(ws, r, cols)
    _ensure_dropdowns(ws, cols)

    if r > 2:  # same look as the row above (hyperlink font on the URL, date formats, ...)
        for col in range(1, max(cols.values()) + 1):
            ws.cell(r, col)._style = copy(ws.cell(r - 1, col)._style)

    for field, value in row.values().items():
        ws.cell(r, cols[field]).value = value or None  # Stage and Remarks: empty
    ws.cell(r, cols["URL"]).hyperlink = row.url
    for field in DATE_FIELDS:
        ws.cell(r, cols[field]).number_format = DATE_FORMAT
    for field in OPTIONAL:
        if field in cols:
            ws.cell(r, cols[field]).value = None

    backup(path, backup_dir)
    _save(wb, path)
    return r


# ---------- pending rows (tracker open in Excel) ----------

def pending_path_for(path) -> Path:
    return Path(path).parent / "pending_rows.json"


def load_pending(pending_path) -> list[TrackerRow]:
    p = Path(pending_path)
    if not p.exists():
        return []
    return [TrackerRow.from_json(d) for d in json.loads(p.read_text(encoding="utf-8"))]


def pending_duplicates(pending_path, url: str | None = None, company: str | None = None,
                       position: str | None = None) -> list[tuple[str, dict]]:
    """Rows still waiting in a pending file that match. -> [(file name, values)]"""
    name = Path(pending_path).name
    rows = [(name, r.values(), [r.url]) for r in load_pending(pending_path)]
    return match_rows(rows, url=url, company=company, position=position)


def _write_pending(pending_path, rows: list[TrackerRow]) -> None:
    p = Path(pending_path)
    if rows:
        p.write_text(json.dumps([r.to_json() for r in rows], indent=2), encoding="utf-8")
    elif p.exists():
        p.unlink()


def save_pending(pending_path, row: TrackerRow) -> None:
    _write_pending(pending_path, load_pending(pending_path) + [row])


def flush_pending(path, pending_path=None, backup_dir=None, append=None) -> list[tuple[int, TrackerRow]]:
    """Write every pending row to the tracker, removing each from the file once it's saved.

    `append(row) -> row number` writes one row; the default is append_row on the Excel file at `path`.
    Raises TrackerLocked (with the unwritten rows still pending) if the tracker is open in Excel;
    any error from `append` likewise leaves the unwritten rows pending.
    """
    pending_path = pending_path or pending_path_for(path)
    append = append or (lambda row: append_row(path, row, backup_dir))
    rows = load_pending(pending_path)
    written = []
    while rows:
        r = append(rows[0])
        written.append((r, rows.pop(0)))
        _write_pending(pending_path, rows)
    return written


# ---------- pending Outcome updates (tracker open in Excel, or the Sheet unreachable) ----------

def pending_updates_path_for(path) -> Path:
    return Path(path).parent / "pending_updates.json"


def load_pending_updates(pending_path) -> list[OutcomeUpdate]:
    p = Path(pending_path)
    if not p.exists():
        return []
    return [OutcomeUpdate.from_json(d) for d in json.loads(p.read_text(encoding="utf-8"))]


def _write_pending_updates(pending_path, updates: list[OutcomeUpdate]) -> None:
    p = Path(pending_path)
    if updates:
        p.write_text(json.dumps([u.to_json() for u in updates], indent=2), encoding="utf-8")
    elif p.exists():
        p.unlink()


def save_pending_updates(pending_path, updates: list[OutcomeUpdate]) -> None:
    _write_pending_updates(pending_path, load_pending_updates(pending_path) + list(updates))


def flush_pending_updates(path, pending_path=None, backup_dir=None, update=None) -> tuple[list, list]:
    """Write every pending update in one go, then empty the pending file. -> (written, not found)

    `update(updates) -> (written, not found)` does the writing; the default is update_outcome on the
    Excel file at `path`. If it raises (TrackerLocked: Excel has the file open), every update stays pending.
    Updates whose job isn't found are reported, not kept: waiting won't make them appear.
    """
    pending_path = pending_path or pending_updates_path_for(path)
    update = update or (lambda updates: update_outcome(path, updates, backup_dir))
    updates = load_pending_updates(pending_path)
    if not updates:
        return [], []
    written, missing = update(updates)
    _write_pending_updates(pending_path, [])
    return written, missing
