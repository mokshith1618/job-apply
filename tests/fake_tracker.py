"""A made-up tracker with the same layout as "Job Application Tracker.xlsx" (see SPEC.md).

Tests build it in a temp folder, so they need no private file and never see real applications.
`headers` puts the columns in another order, to show nothing depends on their position.
"""
from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = ["Company", "Position", "URL", "Posted date/ Source", "Date", "Outcome", "Stage",
           "Last Updated", "Remarks", "Experience", "Key Skills"]
SHUFFLED = ["Remarks", "Outcome", "URL", "Key Skills", "Company", "Last Updated", "Date", "Stage",
            "Position", "Posted date/ Source", "Experience"]
LAST_ROW = 897  # last row with a Company value, so the next row is 898 like the real file
TABLE_END = 1000
LINK_FONT = Font(color="FF0563C1", underline="single")
OUTCOMES = '"Applied,Rejected,Awaiting reply,In progress"'
STAGES = '"Online Assessment,Interview,Initial Screening"'  # the real file's order


def date_check(letter: str) -> str:
    return f'OR(NOT(ISERROR(DATEVALUE({letter}2))), AND(ISNUMBER({letter}2), LEFT(CELL("format", {letter}2))="D"))'


def letter_of(header: str, headers=HEADERS) -> str:
    return get_column_letter(headers.index(header) + 1)


def make_tracker(path: Path, headers=HEADERS) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers)
    col = {h: i for i, h in enumerate(headers, start=1)}
    first_day = datetime(2025, 6, 1)
    for r in range(2, LAST_ROW + 1):
        url = f"https://example.com/jobs/{r}"
        ws.cell(r, col["Company"], f"Company {r}")
        ws.cell(r, col["Position"], f"Role {r}")
        ws.cell(r, col["URL"], url).hyperlink = url
        ws.cell(r, col["URL"]).font = LINK_FONT
        ws.cell(r, col["Posted date/ Source"], "Posted today-Other")
        ws.cell(r, col["Date"], first_day + timedelta(days=r % 400)).number_format = "M/d/yyyy"
        ws.cell(r, col["Outcome"], "Applied")
    ws.freeze_panes = "A2"

    table = Table(displayName="Table1", ref=f"A1:{get_column_letter(len(headers))}{TABLE_END}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    for header, dv in (("Outcome", DataValidation(type="list", formula1=OUTCOMES, allow_blank=True)),
                       ("Stage", DataValidation(type="list", formula1=STAGES, allow_blank=True)),
                       ("Date", None), ("Last Updated", None)):
        letter = letter_of(header, headers)
        dv = dv or DataValidation(type="custom", formula1=date_check(letter), allow_blank=True)
        dv.add(f"{letter}2:{letter}{TABLE_END}")
        ws.add_data_validation(dv)

    wb.save(path)
    return path
