"""Google Sheet tracker against a fake worksheet (no network), plus apply.py runs that write to both."""
import csv
from datetime import date

import pytest
import requests
from openpyxl import load_workbook

import apply
import sheet_tracker
import tracker
from ats.common import JobDetails
from conftest import KNOWN_ROW, KNOWN_URL
from fake_sheet import HEADER, FakeWorksheet
from sheet_tracker import SheetMismatch, SheetUnavailable

TODAY = date(2026, 10, 13)
NEW_URL = "https://job-boards.greenhouse.io/acme/jobs/1234567"
LINKED_URL = "https://jobs.lever.co/linked/abc-123"


def sample_sheet(**kwargs):
    return FakeWorksheet([
        HEADER,
        ["Example Labs", "Software Engineer - Python", "https://jobs.ashbyhq.com/examplelabs/1234abcd",
         "Posted today-Tsenta", "6/20/2026", "Applied"],
        ["Linked Co", "Backend Engineer", "click here", "Posted today-LinkedIn", "7/1/2026", "Rejected"],
        ["", "", "", "", "", "", "", "", "a remark with no company"],
    ], formulas={3: f'=HYPERLINK("{LINKED_URL}", "click here")'}, **kwargs)


def new_row(**overrides):
    fields = dict(company="Acme Robotics", position="Software Engineer I", url=NEW_URL,
                  posted="Posted October 8, 2026-LinkedIn", applied_on=TODAY)
    return tracker.TrackerRow(**{**fields, **overrides})


# ---------- sheet_tracker ----------

def test_row_goes_after_last_company_with_typed_values(tmp_path):
    ws = sample_sheet()
    assert sheet_tracker.append_row(ws, new_row(), tmp_path / "backups") == 4
    assert dict(ws.updates) == {
        "A4": "Acme Robotics", "B4": "Software Engineer I", "C4": f'=HYPERLINK("{NEW_URL}", "{NEW_URL}")',
        "D4": "Posted October 8, 2026-LinkedIn", "E4": "2026-10-13", "F4": "Applied", "H4": "2026-10-13",
    }
    assert {c for c, o in ws.input_options.items() if o == "USER_ENTERED"} == {"C4", "E4", "H4"}


def test_format_and_dropdown_copied_from_row_above_and_date_format_set(tmp_path):
    ws = sample_sheet()
    sheet_tracker.append_row(ws, new_row(), tmp_path)
    copies = [q["copyPaste"] for q in ws.spreadsheet.requests if "copyPaste" in q]
    assert [c["pasteType"] for c in copies] == ["PASTE_FORMAT", "PASTE_DATA_VALIDATION"]
    for c in copies:
        assert (c["source"]["sheetId"], c["source"]["startRowIndex"], c["destination"]["startRowIndex"]) == (7, 2, 3)
        assert c["destination"]["endColumnIndex"] == 11
    cells = [q["repeatCell"] for q in ws.spreadsheet.requests if "repeatCell" in q]
    assert [(c["range"]["startRowIndex"], c["range"]["startColumnIndex"]) for c in cells] == [(3, 4), (3, 7)]  # E, H
    for cell in cells:
        assert cell["cell"]["userEnteredFormat"]["numberFormat"] == {"type": "DATE", "pattern": "M/d/yyyy"}


def test_text_is_written_raw_and_unchanged(tmp_path):
    ws = sample_sheet()
    sheet_tracker.append_row(ws, new_row(company="=IMPORTXML(1)", position="SDE 1-2", posted="+Lead 10/20",
                                         url='https://x.com/a"b'), tmp_path)
    assert ws.values[3][:4] == ["=IMPORTXML(1)", "SDE 1-2", 'https://x.com/a"b', "+Lead 10/20"]
    assert [ws.input_options[c] for c in ("A4", "B4", "D4", "F4")] == ["RAW"] * 4
    assert ws.formulas[4] == '=HYPERLINK("https://x.com/a""b", "https://x.com/a""b")'


def test_wrong_header_writes_nothing(tmp_path):
    ws = FakeWorksheet([["Name", "Email"], ["x", "y"]])
    with pytest.raises(SheetMismatch):
        sheet_tracker.append_row(ws, new_row(), tmp_path / "backups")
    assert ws.updates == [] and ws.spreadsheet.requests == []
    assert not (tmp_path / "backups").exists()


def test_missing_sheet_header_is_named(tmp_path):
    ws = FakeWorksheet([[h for h in HEADER if h != "Last Updated"]])
    with pytest.raises(SheetMismatch, match="Last Updated"):
        sheet_tracker.append_row(ws, new_row(), tmp_path / "backups")
    assert ws.updates == [] and not (tmp_path / "backups").exists()


def test_shuffled_sheet_columns_append_and_read(tmp_path):
    order = ["Remarks", "Outcome", "URL", "Company", "Last Updated", "Date", "Stage", "Position",
             "Posted date/ Source"]
    ws = FakeWorksheet([order, ["", "Applied", "https://x.com/1", "Old Co", "", "6/20/2026", "", "SWE", "LinkedIn"]])
    assert sheet_tracker.append_row(ws, new_row(), tmp_path) == 3
    assert dict(ws.updates) == {"D3": "Acme Robotics", "H3": "Software Engineer I",
                                "C3": f'=HYPERLINK("{NEW_URL}", "{NEW_URL}")', "I3": "Posted October 8, 2026-LinkedIn",
                                "F3": "2026-10-13", "B3": "Applied", "E3": "2026-10-13"}
    rows = sheet_tracker.read_rows(ws)
    assert [(r, v["Company"], v["Position"], urls) for r, v, urls in rows] == [
        (2, "Old Co", "SWE", ["https://x.com/1"]), (3, "Acme Robotics", "Software Engineer I", [NEW_URL, NEW_URL])]


def test_backup_csv_holds_values_before_the_write(tmp_path):
    ws = sample_sheet()
    sheet_tracker.append_row(ws, new_row(), tmp_path / "backups")
    (backup,) = (tmp_path / "backups").glob("sheet-*.csv")
    with open(backup, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[0] == HEADER and len(rows) == 4 and rows[1][0] == "Example Labs"


def test_full_sheet_gets_more_rows(tmp_path):
    ws = sample_sheet(row_count=3)
    assert sheet_tracker.append_row(ws, new_row(), tmp_path) == 4
    assert ws.added_rows == sheet_tracker.SHEET_GROWTH


def test_second_row_goes_below_first(tmp_path):
    ws = sample_sheet()
    sheet_tracker.append_row(ws, new_row(), tmp_path)
    assert sheet_tracker.append_row(ws, new_row(url=NEW_URL + "9"), tmp_path) == 5


def test_duplicates_by_url_hyperlink_target_and_company_position():
    ws = sample_sheet()
    hits = sheet_tracker.find_duplicates(ws, url="https://www.jobs.ashbyhq.com/examplelabs/1234abcd/?utm_source=x")
    assert [r for r, _ in hits] == [2]
    assert hits[0][1]["Company"] == "Example Labs"
    assert [r for r, _ in sheet_tracker.find_duplicates(ws, url=LINKED_URL + "/apply")] == [3]
    assert [r for r, _ in sheet_tracker.find_duplicates(ws, company="linked co", position="BACKEND  engineer")] == [3]
    assert sheet_tracker.find_duplicates(ws, url=NEW_URL, company="Acme", position="SWE") == []


def test_network_errors_become_sheet_unavailable(tmp_path):
    ws = sample_sheet(fail=requests.exceptions.ConnectionError("offline"))
    with pytest.raises(SheetUnavailable, match="offline"):
        sheet_tracker.append_row(ws, new_row(), tmp_path)
    with pytest.raises(SheetUnavailable):
        sheet_tracker.find_duplicates(ws, url=NEW_URL)


def test_missing_key_file_is_sheet_unavailable(tmp_path, monkeypatch):
    monkeypatch.undo()  # the autouse guard replaces open_sheet; test the real one
    with pytest.raises(SheetUnavailable, match="key not found"):
        sheet_tracker.open_sheet("id", "Sheet1", tmp_path / "missing.json")


# ---------- apply.py writes to both ----------

class Script:
    def __init__(self, *answers):
        self.answers = list(answers)
        self.prompts = []

    def __call__(self, prompt=""):
        self.prompts.append(prompt)
        assert self.answers, f"unexpected prompt: {prompt!r}"
        return self.answers.pop(0)


def fetch_ok(url):
    return JobDetails("Acme Robotics", "Software Engineer I", date(2026, 10, 8))


def run(tracker_file, ws, script, url=NEW_URL, extra=()):
    return apply.main([url, "--tracker", str(tracker_file), *extra], input_fn=script, fetch=fetch_ok,
                      today=TODAY, sheet_ws=ws)


def excel(path, ref):
    return load_workbook(path)["Sheet1"][ref].value


def test_row_goes_to_sheet_and_excel(tracker_file, capsys):
    ws = sample_sheet()
    assert run(tracker_file, ws, Script("")) == 0
    assert ws.values[3][0] == "Acme Robotics"
    assert excel(tracker_file, "A898") == "Acme Robotics"
    out = capsys.readouterr().out
    assert "Added row 4 to the Google Sheet" in out and "Added row 898" in out
    assert list((tracker_file.parent / "backups").glob("sheet-*.csv"))


def test_duplicate_found_in_sheet_is_shown_with_the_card(tracker_file, capsys):
    ws = sample_sheet()
    before = tracker_file.read_bytes()
    script = Script("s")
    assert run(tracker_file, ws, script, url=LINKED_URL) == 0
    assert script.prompts == [apply.PAUSE_PROMPT]
    assert "row 3: Linked Co" in capsys.readouterr().out
    assert ws.updates == [] and tracker_file.read_bytes() == before


def test_skip_writes_nowhere(tracker_file):
    ws = sample_sheet()
    before = tracker_file.read_bytes()
    assert run(tracker_file, ws, Script("s")) == 0
    assert ws.updates == [] and tracker_file.read_bytes() == before


def test_sheet_down_at_write_keeps_row_pending_then_flushes(tracker_file, capsys):
    ws = sample_sheet()
    real_update = ws.batch_update

    def offline(*args, **kwargs):
        raise requests.exceptions.ConnectionError("offline")

    ws.batch_update = offline
    assert run(tracker_file, ws, Script("")) == 0
    pending = tracker_file.parent / "pending_sheet.json"
    assert [r.company for r in tracker.load_pending(pending)] == ["Acme Robotics"]
    assert excel(tracker_file, "A898") == "Acme Robotics"  # Excel still written
    assert "pending_sheet.json" in capsys.readouterr().out

    ws.batch_update = real_update  # back online
    assert run(tracker_file, ws, Script("s"), url=NEW_URL + "9") == 0
    assert "pending_sheet.json: Acme Robotics" in capsys.readouterr().out  # same company + position, still waiting
    assert ws.values[3][0] == "" and pending.exists()  # skipped: nothing flushed

    script = Script("")  # applied: the waiting row is written first, then this one
    assert run(tracker_file, ws, script, url=NEW_URL + "9") == 0
    assert script.prompts == [apply.PAUSE_PROMPT]
    assert ws.values[3][0] == "Acme Robotics" and ws.values[4][2] == NEW_URL + "9" and not pending.exists()
    assert "Wrote pending row 4 to the Google Sheet" in capsys.readouterr().out


def test_wrong_sheet_keeps_row_pending(tracker_file):
    ws = FakeWorksheet([["Name", "Email"]])
    assert run(tracker_file, ws, Script("")) == 0
    assert ws.updates == []
    assert tracker.load_pending(tracker_file.parent / "pending_sheet.json")
    assert excel(tracker_file, "A898") == "Acme Robotics"


def test_no_sheet_flag_is_excel_only(tracker_file):
    ws = sample_sheet()
    assert run(tracker_file, ws, Script(""), extra=["--no-sheet"]) == 0
    assert ws.updates == [] and len(ws.values) == 4
    assert excel(tracker_file, "A898") == "Acme Robotics"
    assert not (tracker_file.parent / "pending_sheet.json").exists()


def test_config_sheet_unreachable_falls_back_to_excel(dup_tracker, tmp_path, monkeypatch, capsys):
    tracker_file = dup_tracker
    config = tmp_path / "config.yaml"
    config.write_text(f"sheet_id: abc\nlocal_tracker: '{tracker_file}'\n", encoding="utf-8")

    def unreachable(sheet_id, tab, key_file):
        assert (sheet_id, tab, key_file.name) == ("abc", "Sheet1", "google-key.json")
        raise SheetUnavailable("offline")

    monkeypatch.setattr(sheet_tracker, "open_sheet", unreachable)
    dup_url = KNOWN_URL
    assert apply.main([dup_url, "--config", str(config)], input_fn=Script("s"),
                      fetch=fetch_ok, today=TODAY) == 0  # duplicate found in the Excel file
    out = capsys.readouterr().out
    assert "Google Sheet unavailable: offline" in out and f"row {KNOWN_ROW}" in out

    assert apply.main([NEW_URL, "--config", str(config)], input_fn=Script(""), fetch=fetch_ok, today=TODAY) == 0
    assert excel(tracker_file, "A898") == "Acme Robotics"
    assert tracker.load_pending(tracker_file.parent / "pending_sheet.json")


def test_no_config_is_excel_only(tracker_file):
    assert apply.main([NEW_URL, "--tracker", str(tracker_file)], input_fn=Script(""),
                      fetch=fetch_ok, today=TODAY) == 0
    assert excel(tracker_file, "A898") == "Acme Robotics"
    assert not (tracker_file.parent / "pending_sheet.json").exists()


def test_sheet_without_stage_column_writes_nothing(tmp_path):
    old = ["Company", "Position", "URL", "Posted date/ Source", "Date", "Outcome", "Remarks"]
    ws = FakeWorksheet([old, ["Acme", "SWE", "https://example.com/1", "", "6/20/2026", "Applied", ""]])
    with pytest.raises(SheetMismatch):
        sheet_tracker.append_row(ws, new_row(), tmp_path / "backups")
    assert ws.updates == [] and not (tmp_path / "backups").exists()


def test_excel_missing_header_stops_apply_before_any_write(tracker_file, capsys):
    wb = load_workbook(tracker_file)
    wb["Sheet1"]["I1"] = "Notes"  # Remarks
    wb.save(tracker_file)
    before = tracker_file.read_bytes()
    ws = sample_sheet()
    script = Script()
    assert run(tracker_file, ws, script) == 2
    assert script.prompts == [] and ws.updates == [] and tracker_file.read_bytes() == before
    assert "missing the header Remarks" in capsys.readouterr().out


def test_sheet_header_warning_on_every_run(tracker_file, capsys):
    ws = FakeWorksheet([[h for h in HEADER if h != "Last Updated"]])
    for script, url in ((Script(""), NEW_URL), (Script("s"), NEW_URL + "9"), (Script(""), NEW_URL + "8")):
        assert run(tracker_file, ws, script, url=url) == 0
        out = capsys.readouterr().out
        assert "Warning: the Google Sheet's row 1 is missing the header Last Updated" in out, out
    assert ws.updates == []
    assert len(tracker.load_pending(tracker_file.parent / "pending_sheet.json")) == 2
