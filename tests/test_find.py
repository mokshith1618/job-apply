"""Milestone 7: find.py search and Outcome / Stage update, on the fake tracker and a fake worksheet."""
import json
from datetime import date, datetime

import pytest
import requests
from openpyxl import load_workbook
from rich.console import Console

import find
import sheet_tracker
import tracker
from conftest import KNOWN_COMPANY, KNOWN_POSITION, KNOWN_ROW, KNOWN_URL
from fake_sheet import HEADER, FakeWorksheet
from sheet_tracker import SheetMismatch, SheetUnavailable

ACME_URL = "https://acme.wd5.myworkdayjobs.com/External/job/Austin/Cloud-Engineer_R123"
SP_URL = "https://job-boards.greenhouse.io/spglobal/jobs/555"
LEVER_URL = "https://jobs.lever.co/zenith/abc-123"


def row(r, company, position, url="", outcome="Applied", remarks="", date="10/1/2026", target=None, stage=""):
    values = {"Company": company, "Position": position, "URL": url, "Posted date/ Source": "Posted today-Other",
              "Date": date, "Outcome": outcome, "Stage": stage, "Last Updated": "", "Remarks": remarks}
    urls = [u for u in (url, target) if u]
    return (r, values, urls)


SAMPLE = [
    row(2, "Acme Inc.", "Cloud Engineer", ACME_URL),
    row(3, "S&P Global", "Software Engineer", SP_URL),
    row(4, "Zenith Labs", "Acme Integrations Developer", "click here", target=LEVER_URL),
    row(5, "Other Co", "Data Engineer", "https://example.com/jobs/5"),
]


def found(text, rows=SAMPLE):
    return [r for r, _, _ in tracker.search_rows(rows, text)]


# ---------- search ----------

def test_search_key_cleans_case_punctuation_and_suffixes():
    assert tracker.search_key("  ACME, Inc. ") == "acme"
    assert tracker.search_key("S&P Global") == "s p global"
    assert tracker.search_key("Foo Bar LLC") == "foo bar"
    assert tracker.search_key("Co") == "co"  # a name that is only a suffix stays
    assert tracker.search_key("Costco Wholesale Corp") == "costco wholesale"


def test_match_by_company_ignoring_case_and_suffix():
    assert found("acme") == [2, 4]  # row 4: "Acme" in the position
    assert found("ACME INC") == [2, 4]
    assert found("acme, inc.") == [2, 4]


def test_match_by_position():
    assert found("data engineer") == [5]
    assert found("integrations") == [4]


def test_match_with_punctuation():
    assert found("s&p global") == [3]
    assert found("S & P") == [3]


def test_match_by_url_text_and_hyperlink_target():
    assert found("myworkdayjobs") == [2]
    assert found("acme.wd5") == [2]
    assert found("jobs.lever.co/zenith") == [4]  # only in the hyperlink target
    assert found("spglobal") == [3]


def test_no_match():
    assert found("globex") == []
    assert found("   ") == []


def test_excel_read_rows_has_remarks_and_hyperlink_targets(tracker_file):
    wb = load_workbook(tracker_file)
    ws = wb["Sheet1"]
    ws.cell(10, 1, "Acme Inc.")
    ws.cell(10, 3, "click here").hyperlink = ACME_URL
    ws.cell(10, 7, "Interview")
    ws.cell(10, 9, "Recruiter call")
    for col in (1, 3):  # row 11: no Company, no URL (ws.cell(r, c, None) would leave the value)
        ws.cell(11, col).value = None
    ws.cell(11, 3).hyperlink = None
    wb.save(tracker_file)

    rows = tracker.read_rows(tracker_file)
    by_row = {r: (values, urls) for r, values, urls in rows}
    values, urls = by_row[10]
    assert (values["Company"], values["Stage"], values["Remarks"]) == ("Acme Inc.", "Interview", "Recruiter call")
    assert urls == ["click here", ACME_URL]
    assert isinstance(values["Date"], datetime)
    assert 11 not in by_row  # no Company and no URL: not a job row
    assert [r for r, _, _ in tracker.search_rows(rows, "acme")] == [10]


def test_sheet_read_rows_has_remarks_and_hyperlink_targets():
    ws = FakeWorksheet([
        HEADER,
        ["Acme Inc.", "Cloud Engineer", "click here", "Posted today-Other", "10/1/2026", "Awaiting reply",
         "Online Assessment", "", "OA sent"],
        ["", "", "", "", "", "", "", "", "a remark with no company"],
        ["S&P Global", "Software Engineer", SP_URL, "Posted today-Other", "9/1/2026", "Rejected"],
    ], formulas={2: f'=HYPERLINK("{ACME_URL}", "click here")'})
    rows = sheet_tracker.read_rows(ws)
    assert [r for r, _, _ in rows] == [2, 4]
    assert (rows[0][1]["Stage"], rows[0][1]["Remarks"]) == ("Online Assessment", "OA sent")
    assert rows[0][2] == ["click here", ACME_URL]
    assert (rows[1][1]["Stage"], rows[1][1]["Remarks"]) == ("", "")  # short rows are padded
    assert [r for r, _, _ in tracker.search_rows(rows, "myworkdayjobs")] == [2]


# ---------- selection, status, note, table (find.py) ----------

def test_parse_selection():
    assert find.parse_selection("1,3", 5) == [1, 3]
    assert find.parse_selection("2-4", 5) == [2, 3, 4]
    assert find.parse_selection(" 3 , 1-2, 3 ", 5) == [1, 2, 3]
    assert find.parse_selection("all", 3) == [1, 2, 3]
    assert find.parse_selection("ALL", 2) == [1, 2]
    for bad in ("0", "6", "4-2", "1-9", "x", "1,,2", "1-", "-2", "1.5", "all,1"):
        assert find.parse_selection(bad, 5) is None, bad


def test_status_menu_lists_outcomes_then_stages():
    assert find.STATUS_PROMPT == ("New status: 1 Applied  2 Rejected  3 Awaiting reply  4 Online Assessment  "
                                  "5 Initial Screening  6 Interview ")


def test_parse_status_outcome_keeps_stage_and_stage_sets_awaiting_reply():
    assert find.parse_status("2") == ("Rejected", "")
    assert find.parse_status("awaiting  REPLY") == ("Awaiting reply", "")
    assert find.parse_status("4") == ("Awaiting reply", "Online Assessment")
    assert find.parse_status("5") == ("Awaiting reply", "Initial Screening")
    assert find.parse_status(" online assessment ") == ("Awaiting reply", "Online Assessment")
    assert find.parse_status("INTERVIEW") == ("Awaiting reply", "Interview")
    for bad in ("", "0", "7", "offer", "rej"):
        assert find.parse_status(bad) is None, bad


def test_format_note_appends_with_date():
    day = date(2026, 10, 21)
    assert tracker.format_note(day, "Phone screen", "") == "10/21/2026 Phone screen"
    assert tracker.format_note(day, "Phone screen", None) == "10/21/2026 Phone screen"
    assert tracker.format_note(day, " Phone screen ", "Same day rejection") == \
        "Same day rejection | 10/21/2026 Phone screen"
    assert tracker.format_note(day, "  ", "kept") == "kept"


def test_rows_grouped_by_company_newest_first():
    rows = [
        row(10, "Acme", "Old role", date="8/21/2026"),
        row(11, "Beta", "Newest overall", date="10/20/2026"),
        row(12, "Acme Inc.", "New role", date=datetime(2026, 10, 18)),
        row(13, "acme", "Middle role", date="10/2/2026"),
        row(14, "Acme", "No date", date=""),
        row(15, "Gamma", "Only undated", date="n/a"),
    ]
    assert [r for r, _, _ in find.order_rows(rows)] == [11, 12, 13, 10, 14, 15]


def show(hits, text="acme"):
    console = Console(record=True, width=200, color_system=None)
    shown = find.show_table(console, hits, text)
    return shown, console.export_text()


def test_table_shows_rows_and_domain():
    hits = [
        row(640, "Acme", "DevOps Engineer", SP_URL, outcome="Rejected", remarks="Same day rejection",
            date="8/21/2026", stage="Interview"),
        row(912, "Acme", "Software Engineer I", ACME_URL, date=datetime(2026, 10, 18)),
    ]
    shown, out = show(hits)
    assert [r for r, _, _ in shown] == [912, 640]
    lines = out.splitlines()
    assert lines[0] == '2 jobs found for "acme"'
    first = next(line for line in lines if "912" in line).split()
    assert first == ["1", "912", "10/18/2026", "Acme", "Software", "Engineer", "I", "Applied",
                     "acme.wd5.myworkdayjobs.com"]
    second = next(line for line in lines if "640" in line)
    assert "Rejected  Interview" in second
    assert second.index("Interview") < second.index("Same day rejection") < second.index("job-boards.greenhouse.io")
    assert lines[2].split() == ["#", "Row", "Date", "Company", "Position", "Outcome", "Stage", "Remarks", "Link"]
    assert second.split()[:3] == ["2", "640", "08/21/2026"]


def test_table_link_text_without_url_and_one_job_wording():
    shown, out = show([row(4, "Zenith Labs", "Developer", "click here")], "zenith")
    assert out.splitlines()[0] == '1 job found for "zenith"'
    assert "click here" in out


def test_table_capped_at_50():
    hits = [row(r, "Acme", f"Role {r}") for r in range(2, 57)]
    shown, out = show(hits)
    assert len(shown) == 50
    assert "5 more not shown; narrow the search." in out


def test_no_matches_message():
    shown, out = show([], "acme")
    assert shown == [] and out.strip() == 'No jobs found for "acme".'


# ---------- Excel update ----------

NOTE_DAY = date(2026, 10, 21)


def update(row_hint=KNOWN_ROW, outcome="Rejected", note="", urls=(KNOWN_URL,), company=KNOWN_COMPANY,
           position=KNOWN_POSITION, applied="", stage=""):
    return tracker.OutcomeUpdate(company=company, position=position, urls=list(urls), applied=applied,
                                 row=row_hint, outcome=outcome, note=note, noted_on=NOTE_DAY, stage=stage)


def cells(path, *refs):
    ws = load_workbook(path)["Sheet1"]
    return [ws[ref].value for ref in refs]


def test_outcome_written_and_remark_appended_not_overwritten(dup_tracker):
    wb = load_workbook(dup_tracker)
    wb["Sheet1"].cell(KNOWN_ROW, 9, "Recruiter call")
    wb.save(dup_tracker)
    written, missing = tracker.update_outcome(dup_tracker, [update(note="Phone screen")],
                                              dup_tracker.parent / "backups")
    assert [r for r, _ in written] == [KNOWN_ROW] and missing == []
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"G{KNOWN_ROW}", f"I{KNOWN_ROW}") == \
        ["Rejected", None, "Recruiter call | 10/21/2026 Phone screen"]
    assert len(list((dup_tracker.parent / "backups").glob("tracker-*.xlsx"))) == 1


def test_stage_written_and_no_note_leaves_remarks_alone(dup_tracker):
    tracker.update_outcome(dup_tracker, [update(outcome="Awaiting reply", stage="Online Assessment")],
                           dup_tracker.parent / "b")
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"G{KNOWN_ROW}", f"I{KNOWN_ROW}") == \
        ["Awaiting reply", "Online Assessment", None]


def test_outcome_without_stage_keeps_the_stage(dup_tracker):
    tracker.update_outcome(dup_tracker, [update(outcome="Awaiting reply", stage="Interview")], dup_tracker.parent / "b")
    tracker.update_outcome(dup_tracker, [update(outcome="Rejected")], dup_tracker.parent / "b")
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"G{KNOWN_ROW}") == ["Rejected", "Interview"]


def dropdowns(path):
    ws = load_workbook(path)["Sheet1"]
    return {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation if dv.type == "list"}


TRACKER_DROPDOWNS = {"F2:F1000": '"Applied,Rejected,Awaiting reply,In progress"',
                     "G2:G1000": '"Online Assessment,Interview,Initial Screening"'}


def test_dropdowns_kept_as_they_are_when_complete(dup_tracker, tracker_file):
    assert dropdowns(dup_tracker) == TRACKER_DROPDOWNS
    tracker.update_outcome(dup_tracker, [update(outcome="Awaiting reply", stage="Interview")], dup_tracker.parent / "b")
    assert dropdowns(dup_tracker) == TRACKER_DROPDOWNS  # old values (Online Assessment in F) never come back
    tracker.append_row(tracker_file, tracker.TrackerRow("Acme", "SWE", "https://x.com/1", "Posted today-Other",
                                                        NOTE_DAY), tracker_file.parent / "b")
    assert dropdowns(tracker_file) == TRACKER_DROPDOWNS


def test_missing_dropdown_values_added(dup_tracker):
    wb = load_workbook(dup_tracker)
    for dv in wb["Sheet1"].data_validations.dataValidation:
        if dv.type == "list":
            dv.formula1 = '"Applied,Rejected"' if str(dv.sqref).startswith("F") else '"Online Assessment"'
    wb.save(dup_tracker)
    tracker.update_outcome(dup_tracker, [update()], dup_tracker.parent / "b")
    assert dropdowns(dup_tracker) == {"F2:F1000": '"' + ",".join(tracker.OUTCOMES) + '"',
                                      "G2:G1000": '"Online Assessment,Initial Screening,Interview"'}


def test_pending_update_saved_before_stage_existed_becomes_a_stage(tmp_path):
    old = update(outcome="Interview").to_json()
    del old["stage"]
    pending = tmp_path / "pending_updates.json"
    pending.write_text(json.dumps([old, {**old, "outcome": "Rejected"}]), encoding="utf-8")
    assert [(u.outcome, u.stage) for u in tracker.load_pending_updates(pending)] == \
        [("Awaiting reply", "Interview"), ("Rejected", "")]


def test_row_found_again_when_the_hint_is_stale(dup_tracker):
    written, missing = tracker.update_outcome(dup_tracker, [update(row_hint=300)], dup_tracker.parent / "b")
    assert [r for r, _ in written] == [KNOWN_ROW] and missing == []
    assert cells(dup_tracker, "F300", f"F{KNOWN_ROW}") == ["Applied", "Rejected"]


def test_row_without_url_found_by_company_position_and_date(tracker_file):
    wb = load_workbook(tracker_file)
    ws = wb["Sheet1"]
    for r, day in ((50, datetime(2026, 9, 1)), (60, datetime(2026, 10, 1))):
        ws.cell(r, 1, "Zenith")
        ws.cell(r, 2, "Developer")
        ws.cell(r, 3).value = None
        ws.cell(r, 3).hyperlink = None
        ws.cell(r, 5, day)
    wb.save(tracker_file)
    u = update(row_hint=2, urls=(), company="Zenith", position="Developer", applied="2026-10-01")
    written, missing = tracker.update_outcome(tracker_file, [u], tracker_file.parent / "b")
    assert [r for r, _ in written] == [60]
    assert cells(tracker_file, "F50", "F60") == ["Applied", "Rejected"]


def test_ambiguous_or_missing_row_not_written(dup_tracker):
    wb = load_workbook(dup_tracker)
    ws = wb["Sheet1"]
    ws.cell(KNOWN_ROW + 1, 1, KNOWN_COMPANY)
    ws.cell(KNOWN_ROW + 1, 2, KNOWN_POSITION)
    ws.cell(KNOWN_ROW + 1, 3, KNOWN_URL).hyperlink = KNOWN_URL
    ws.cell(KNOWN_ROW + 1, 5, ws.cell(KNOWN_ROW, 5).value)
    wb.save(dup_tracker)
    before = dup_tracker.read_bytes()
    gone = update(urls=("https://jobs.lever.co/nobody/1",), company="Nobody", position="None")
    twins = update(row_hint=5)  # two rows with this job and no way to tell them apart
    written, missing = tracker.update_outcome(dup_tracker, [gone, twins], dup_tracker.parent / "backups")
    assert written == [] and missing == [gone, twins]
    assert dup_tracker.read_bytes() == before and not (dup_tracker.parent / "backups").exists()


def test_locked_excel_keeps_update_pending_until_next_flush(dup_tracker, monkeypatch):
    pending = tracker.pending_updates_path_for(dup_tracker)
    tracker.save_pending_updates(pending, [update(note="Phone screen")])

    def locked(wb, path):
        raise tracker.TrackerLocked("open in Excel")

    real_save = tracker._save
    monkeypatch.setattr(tracker, "_save", locked)
    with pytest.raises(tracker.TrackerLocked):
        tracker.flush_pending_updates(dup_tracker, backup_dir=dup_tracker.parent / "b")
    assert tracker.load_pending_updates(pending) == [update(note="Phone screen")]
    assert cells(dup_tracker, f"F{KNOWN_ROW}") == ["Applied"]

    monkeypatch.setattr(tracker, "_save", real_save)
    written, missing = tracker.flush_pending_updates(dup_tracker, backup_dir=dup_tracker.parent / "b")
    assert [r for r, _ in written] == [KNOWN_ROW] and not pending.exists()
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"I{KNOWN_ROW}") == ["Rejected", "10/21/2026 Phone screen"]


# ---------- Sheet update ----------

LIST_RULE = {"condition": {"type": "ONE_OF_LIST", "values": [
    {"userEnteredValue": "Applied"}, {"userEnteredValue": "Rejected"}]},
    "strict": True, "showCustomUi": True}
STAGE_RULE = {"condition": {"type": "ONE_OF_LIST", "values": [  # the real tracker's order
    {"userEnteredValue": "Online Assessment"}, {"userEnteredValue": "Interview"},
    {"userEnteredValue": "Initial Screening"}]}, "strict": True, "showCustomUi": True}


def sheet(**kwargs):
    return FakeWorksheet([
        HEADER,
        ["Acme Inc.", "Cloud Engineer", "click here", "Posted today-Other", "10/1/2026", "Applied", "", "", "OA sent"],
        ["S&P Global", "Software Engineer", SP_URL, "Posted today-Other", "9/1/2026", "Applied"],
    ], formulas={2: f'=HYPERLINK("{ACME_URL}", "click here")'}, **kwargs)


def at(ws, r, *fields):
    """Values of row `r` in the fake worksheet, by header name."""
    header = ws.values[0]
    row = ws.values[r - 1] + [""] * len(header)
    return [row[header.index(f)] for f in fields]


def sheet_update(ws, r, **kwargs):
    (found_row,) = [x for x in sheet_tracker.read_rows(ws) if x[0] == r]
    return tracker.OutcomeUpdate.for_row(*found_row, outcome=kwargs.get("outcome", "Rejected"),
                                         note=kwargs.get("note", ""), noted_on=NOTE_DAY,
                                         stage=kwargs.get("stage", ""))


def test_sheet_outcome_and_remark_written(tmp_path):
    ws = sheet()
    updates = [sheet_update(ws, 2, note="=Phone screen"),
               sheet_update(ws, 3, outcome="Awaiting reply", stage="Interview")]
    written, missing = sheet_tracker.update_outcome(ws, updates, tmp_path / "backups")
    assert [r for r, _ in written] == [2, 3] and missing == []
    assert at(ws, 2, "Outcome", "Stage", "Remarks") == ["Rejected", "", "OA sent | 10/21/2026 =Phone screen"]
    assert at(ws, 3, "Outcome", "Stage", "Remarks") == ["Awaiting reply", "Interview", ""]
    assert len(list((tmp_path / "backups").glob("sheet-*.csv"))) == 1


def test_sheet_rejected_keeps_stage_and_remark_that_looks_like_a_formula_is_kept_as_text(tmp_path):
    ws = sheet()
    ws.values[1][6], ws.values[1][8] = "Interview", "=SUM(1)"
    sheet_tracker.update_outcome(ws, [sheet_update(ws, 2, note="10/20")], tmp_path)
    assert at(ws, 2, "Outcome", "Stage", "Remarks") == ["Rejected", "Interview", "=SUM(1) | 10/21/2026 10/20"]
    assert (ws.input_options["F2"], ws.input_options["I2"]) == ("RAW", "RAW")
    assert "G2" not in ws.input_options  # Stage kept: not written


def test_sheet_row_found_again_after_a_shift(tmp_path):
    ws = sheet()
    u = sheet_update(ws, 3)
    ws.values.insert(1, ["New Co", "New role", "https://x.com/9", "", "10/20/2026", "Applied"])
    written, _ = sheet_tracker.update_outcome(ws, [u], tmp_path)
    assert [r for r, _ in written] == [4] and at(ws, 4, "Outcome") == ["Rejected"]


def test_sheet_dropdown_gets_missing_outcome(tmp_path):
    ws = sheet(validation=LIST_RULE, stage_validation=STAGE_RULE)
    sheet_tracker.update_outcome(ws, [sheet_update(ws, 2)], tmp_path)
    assert [p["ranges"] for p in ws.spreadsheet.metadata_params] == ["'Sheet1'!F2", "'Sheet1'!G2"]
    (req,) = [q["setDataValidation"] for q in ws.spreadsheet.requests]  # the Stage list was complete
    assert req["range"] == {"sheetId": 7, "startRowIndex": 1, "endRowIndex": 1000,
                            "startColumnIndex": 5, "endColumnIndex": 6}
    assert [v["userEnteredValue"] for v in req["rule"]["condition"]["values"]] == list(tracker.OUTCOMES)
    assert req["rule"]["strict"] is True and req["rule"]["showCustomUi"] is True


def test_sheet_dropdown_left_alone_when_complete_missing_or_not_a_list(tmp_path, capsys):
    complete = {**LIST_RULE, "condition": {"type": "ONE_OF_LIST", "values": [
        {"userEnteredValue": o} for o in tracker.OUTCOMES]}}
    by_range = {"condition": {"type": "ONE_OF_RANGE", "values": [{"userEnteredValue": "=Lists!A1:A3"}]}}
    for rule in (complete, None, by_range):
        ws = sheet(validation=rule, stage_validation=STAGE_RULE)
        sheet_tracker.update_outcome(ws, [sheet_update(ws, 2)], tmp_path)
        assert ws.spreadsheet.requests == []
    assert f"Outcome dropdown is a ONE_OF_RANGE rule; check it offers {', '.join(tracker.OUTCOMES)}" \
        in capsys.readouterr().out


def test_sheet_stage_dropdown_gets_missing_stage(tmp_path):
    two = {**STAGE_RULE, "condition": {"type": "ONE_OF_LIST", "values": STAGE_RULE["condition"]["values"][:2]}}
    ws = sheet(stage_validation=two)
    sheet_tracker.update_outcome(ws, [sheet_update(ws, 2)], tmp_path)
    (req,) = [q["setDataValidation"] for q in ws.spreadsheet.requests]
    assert (req["range"]["startColumnIndex"], req["range"]["endColumnIndex"]) == (6, 7)
    assert [v["userEnteredValue"] for v in req["rule"]["condition"]["values"]] == \
        ["Online Assessment", "Interview", "Initial Screening"]


def test_sheet_missing_row_and_wrong_header_write_nothing(tmp_path):
    ws = sheet()
    gone = update(urls=("https://jobs.lever.co/nobody/1",), company="Nobody", position="None")
    assert sheet_tracker.update_outcome(ws, [gone], tmp_path / "b") == ([], [gone])
    assert ws.updates == [] and not (tmp_path / "b").exists()
    ws = FakeWorksheet([["Name", "Email"], ["x", "y"]])
    with pytest.raises(SheetMismatch):
        sheet_tracker.update_outcome(ws, [gone], tmp_path / "b")
    assert ws.updates == []


def test_sheet_offline_is_sheet_unavailable(tmp_path):
    ws = sheet()
    u = sheet_update(ws, 2)
    ws.fail = requests.exceptions.ConnectionError("offline")
    with pytest.raises(SheetUnavailable):
        sheet_tracker.update_outcome(ws, [u], tmp_path)


# ---------- find.py runs (fake tracker + fake worksheet) ----------

class Script:
    def __init__(self, *answers):
        self.answers = list(answers)
        self.prompts = []

    def __call__(self, prompt=""):
        self.prompts.append(prompt)
        assert self.answers, f"unexpected prompt: {prompt!r}"
        return self.answers.pop(0)


def known_sheet(**kwargs):
    return FakeWorksheet([
        HEADER,
        [KNOWN_COMPANY, KNOWN_POSITION, KNOWN_URL, "Posted today-Tsenta", "6/20/2026", "Applied"],
        ["Other Co", "Data Engineer", "https://example.com/other", "Posted today-Other", "7/1/2026", "Applied"],
    ], **kwargs)


def run_find(tracker_file, ws, script, text="example labs", extra=()):
    return find.main([*text.split(), "--tracker", str(tracker_file), *extra], input_fn=script,
                     sheet_ws=ws, today=NOTE_DAY)


def test_update_goes_to_sheet_then_excel_and_table_shown_again(dup_tracker, capsys):
    ws = known_sheet()
    script = Script("1", "2", "Phone screen", "y", "")
    assert run_find(dup_tracker, ws, script) == 0
    assert script.prompts == [find.ROWS_PROMPT, find.STATUS_PROMPT, find.NOTE_PROMPT,
                              "Update 1 row to Rejected? [y/N] ", find.ROWS_PROMPT]
    assert at(ws, 2, "Outcome", "Stage", "Remarks") == ["Rejected", "", "10/21/2026 Phone screen"]
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"G{KNOWN_ROW}", f"I{KNOWN_ROW}") == \
        ["Rejected", None, "10/21/2026 Phone screen"]
    out = capsys.readouterr().out
    assert out.index("Updated row 2 in the Google Sheet") < out.index(f"Updated row {KNOWN_ROW} in tracker.xlsx")
    assert out.count('1 job found for "example labs"') == 2
    assert list((dup_tracker.parent / "backups").glob("sheet-*.csv"))
    assert list((dup_tracker.parent / "backups").glob("tracker-*.xlsx"))


def test_invalid_answers_are_asked_again(dup_tracker):
    ws = known_sheet()
    script = Script("x", "5", "1", "offer", "7", "4", "", "y", "")
    assert run_find(dup_tracker, ws, script) == 0
    assert script.prompts[:6] == [find.ROWS_PROMPT] * 3 + [find.STATUS_PROMPT] * 3
    assert script.prompts[7] == "Update 1 row to Online Assessment? [y/N] "
    assert at(ws, 2, "Outcome", "Stage") == ["Awaiting reply", "Online Assessment"]


def test_interview_then_rejected_keeps_the_stage(dup_tracker):
    ws = known_sheet()
    assert run_find(dup_tracker, ws, Script("1", "6", "", "y", "1", "rejected", "", "y", "")) == 0
    assert at(ws, 2, "Outcome", "Stage") == ["Rejected", "Interview"]
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"G{KNOWN_ROW}") == ["Rejected", "Interview"]


def test_anything_but_y_writes_nothing(dup_tracker, capsys):
    ws = known_sheet()
    before = dup_tracker.read_bytes()
    assert run_find(dup_tracker, ws, Script("1", "2", "note", "n", "")) == 0
    assert ws.updates == [] and dup_tracker.read_bytes() == before
    assert "Nothing changed." in capsys.readouterr().out


def test_no_matches(dup_tracker, capsys):
    script = Script()
    assert run_find(dup_tracker, known_sheet(), script, text="globex") == 0
    assert script.prompts == []
    assert 'No jobs found for "globex".' in capsys.readouterr().out


def test_excel_search_when_no_sheet_capped_at_50(dup_tracker, capsys):
    script = Script("51", "")
    assert run_find(dup_tracker, known_sheet(), script, text="company", extra=["--no-sheet"]) == 0
    out = capsys.readouterr().out
    assert '895 jobs found for "company"' in out and "845 more not shown; narrow the search." in out
    assert script.prompts == [find.ROWS_PROMPT] * 2


def test_sheet_down_at_write_keeps_update_pending_then_flushes(dup_tracker, capsys):
    ws = known_sheet()
    real_update = ws.batch_update

    def offline(*args, **kwargs):
        raise requests.exceptions.ConnectionError("offline")

    ws.batch_update = offline
    assert run_find(dup_tracker, ws, Script("1", "2", "", "y", "")) == 0
    pending = dup_tracker.parent / "pending_sheet_updates.json"
    assert [u.outcome for u in tracker.load_pending_updates(pending)] == ["Rejected"]
    assert cells(dup_tracker, f"F{KNOWN_ROW}") == ["Rejected"]  # Excel still written
    assert "pending_sheet_updates.json" in capsys.readouterr().out

    ws.batch_update = real_update  # back online: the waiting update goes first, then the new one
    assert run_find(dup_tracker, ws, Script("1", "4", "", "y", "")) == 0
    assert [(c, v) for c, v in ws.updates if c in ("F2", "G2")] == \
        [("F2", "Rejected"), ("F2", "Awaiting reply"), ("G2", "Online Assessment")]
    assert not pending.exists()


def test_sheet_unreachable_at_start_searches_excel(dup_tracker, tmp_path, monkeypatch, capsys):
    config = tmp_path / "config.yaml"
    config.write_text(f"sheet_id: abc\nlocal_tracker: '{dup_tracker}'\n", encoding="utf-8")

    def unreachable(sheet_id, tab, key_file):
        raise SheetUnavailable("offline")

    monkeypatch.setattr(sheet_tracker, "open_sheet", unreachable)
    assert find.main(["example", "labs", "--config", str(config)], input_fn=Script("1", "6", "", "y", ""),
                     today=NOTE_DAY) == 0
    out = capsys.readouterr().out
    assert "Google Sheet unavailable: offline" in out and "Searching tracker.xlsx instead" in out
    assert cells(dup_tracker, f"F{KNOWN_ROW}", f"G{KNOWN_ROW}") == ["Awaiting reply", "Interview"]
    assert tracker.load_pending_updates(dup_tracker.parent / "pending_sheet_updates.json")


def locked_once(monkeypatch, times=1):
    real_save = tracker._save
    calls = {"n": 0}

    def save(wb, path):
        calls["n"] += 1
        if calls["n"] <= times:
            raise tracker.TrackerLocked("open in Excel")
        real_save(wb, path)

    monkeypatch.setattr(tracker, "_save", save)
    return real_save


def test_excel_open_asks_to_close_it_then_writes(dup_tracker, monkeypatch):
    locked_once(monkeypatch)
    script = Script("1", "2", "", "y", "", "")
    assert run_find(dup_tracker, None, script, extra=["--no-sheet"]) == 0
    assert "Close Excel, then press Enter to retry, s = leave it pending" in script.prompts[4]
    assert cells(dup_tracker, f"F{KNOWN_ROW}") == ["Rejected"]
    assert not tracker.pending_updates_path_for(dup_tracker).exists()


def test_excel_open_s_keeps_update_pending_and_apply_writes_it_later(dup_tracker, monkeypatch, capsys):
    from test_sheet_tracker import fetch_ok

    import apply

    real_save = locked_once(monkeypatch, times=99)
    assert run_find(dup_tracker, None, Script("1", "2", "Same day rejection", "y", "s", ""),
                    extra=["--no-sheet"]) == 0
    pending = tracker.pending_updates_path_for(dup_tracker)
    assert tracker.load_pending_updates(pending)[0].note == "Same day rejection"
    assert cells(dup_tracker, f"F{KNOWN_ROW}") == ["Applied"]
    assert "Update kept in pending_updates.json" in capsys.readouterr().out

    monkeypatch.setattr(tracker, "_save", real_save)  # Excel closed; apply.py's next write also writes it
    assert apply.main(["https://job-boards.greenhouse.io/acme/jobs/1", "--tracker", str(dup_tracker),
                       "--no-sheet"], input_fn=Script(""), fetch=fetch_ok, today=NOTE_DAY) == 0
    assert cells(dup_tracker, "A898", f"F{KNOWN_ROW}", f"I{KNOWN_ROW}") == \
        ["Acme Robotics", "Rejected", "10/21/2026 Same day rejection"]
    assert not pending.exists()


def test_end_of_input_or_ctrl_c_quits_without_writing(dup_tracker, capsys):
    before = dup_tracker.read_bytes()
    for stop in (EOFError, KeyboardInterrupt):
        def script(prompt="", stop=stop):
            raise stop
        assert run_find(dup_tracker, None, script, extra=["--no-sheet"]) == 0
    assert dup_tracker.read_bytes() == before
    assert capsys.readouterr().out.count("Stopped.") == 2


def test_byte_order_mark_from_a_pipe_is_ignored(dup_tracker):
    script = Script("﻿")
    assert run_find(dup_tracker, None, script, extra=["--no-sheet"]) == 0
    assert script.prompts == [find.ROWS_PROMPT]


def test_excel_missing_header_stops_find(dup_tracker, capsys):
    wb = load_workbook(dup_tracker)
    wb["Sheet1"]["G1"] = "Status"  # Stage
    wb.save(dup_tracker)
    before = dup_tracker.read_bytes()
    script = Script()
    assert run_find(dup_tracker, known_sheet(), script) == 2
    assert script.prompts == [] and dup_tracker.read_bytes() == before
    assert "tracker.xlsx: row 1 is missing the header Stage" in capsys.readouterr().out
