import hashlib
import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TRACKER_COPY = ROOT / "tracker_copy.xlsx"
REAL_TRACKER = ROOT / "Job Application Tracker.xlsx"


def _digest(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


@pytest.fixture(scope="session", autouse=True)
def originals_untouched():
    """Tests use a generated fake tracker; tracker_copy.xlsx and the real tracker must never change."""
    before = (_digest(TRACKER_COPY), _digest(REAL_TRACKER))
    yield
    assert (_digest(TRACKER_COPY), _digest(REAL_TRACKER)) == before


@pytest.fixture(scope="session")
def fake_tracker_template(tmp_path_factory) -> Path:
    """Built once per run: a made-up tracker with the real file's layout (tests/fake_tracker.py)."""
    from fake_tracker import make_tracker

    return make_tracker(tmp_path_factory.mktemp("template") / "tracker.xlsx")


@pytest.fixture
def tracker_file(tmp_path, fake_tracker_template) -> Path:
    """A fresh copy of the fake tracker for each test (never tracker_copy.xlsx or the real tracker)."""
    dest = tmp_path / "tracker.xlsx"
    shutil.copy2(fake_tracker_template, dest)
    return dest


@pytest.fixture(autouse=True)
def no_real_google_sheet(tmp_path, monkeypatch):
    """Tests never read the real config.yaml or open the real Google Sheet; they pass a FakeWorksheet."""
    import apply
    import sheet_tracker

    monkeypatch.setattr(apply, "DEFAULT_CONFIG", tmp_path / "no-config.yaml")

    def refuse(*args, **kwargs):
        raise AssertionError("tests must not open the real Google Sheet")

    monkeypatch.setattr(sheet_tracker, "open_sheet", refuse)


@pytest.fixture(autouse=True)
def no_real_browser(tmp_path, monkeypatch):
    """apply.main never opens Chrome or reads the real profile.yaml/answers.yaml; tests pass fill= instead."""
    import apply

    monkeypatch.setattr(apply, "FILLERS", {})
    monkeypatch.setattr(apply, "DEFAULT_PROFILE", tmp_path / "no-profile.yaml")
    monkeypatch.setattr(apply, "DEFAULT_ANSWERS", tmp_path / "no-answers.yaml")
    monkeypatch.setattr(apply, "DEFAULT_COMPANIES", tmp_path / "companies.yaml")
    monkeypatch.setattr(apply, "BROWSER_PROFILE", tmp_path / "no-browser-profile")


@pytest.fixture(autouse=True)
def test_keywords(tmp_path, monkeypatch):
    """apply.main reads a made-up matching.yaml, not mine."""
    import apply

    path = tmp_path / "matching.yaml"
    path.write_text("keywords:\n  sponsorship: [sponsor, visa, OPT]\n  citizenship: [citizen, clearance]\n"
                    "  citizenship_ignore: [regardless of]\n", encoding="utf-8")
    monkeypatch.setattr(apply, "DEFAULT_MATCHING", path)


# A known job planted at row 199 of the fake tracker, for the duplicate tests.
KNOWN_ROW = 199
KNOWN_COMPANY = "Example Labs, Inc."
KNOWN_POSITION = "Software Engineer - Python"
KNOWN_URL = "https://jobs.ashbyhq.com/examplelabs/00000000-0000-4000-8000-000000000199"


@pytest.fixture
def dup_tracker(tracker_file) -> Path:
    """tracker_file with the KNOWN_* job in row 199 (C as text and hyperlink, like real rows)."""
    from openpyxl import load_workbook

    wb = load_workbook(tracker_file)
    ws = wb["Sheet1"]
    ws.cell(KNOWN_ROW, 1, KNOWN_COMPANY)
    ws.cell(KNOWN_ROW, 2, KNOWN_POSITION)
    ws.cell(KNOWN_ROW, 3, KNOWN_URL).hyperlink = KNOWN_URL
    wb.save(tracker_file)
    return tracker_file
