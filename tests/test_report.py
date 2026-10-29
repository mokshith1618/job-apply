"""Milestone 9: report.py, a read-only page of the tracker, on the fake tracker and a fake worksheet."""
import hashlib
import re
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from fake_sheet import HEADER, FakeWorksheet
from fake_tracker import SHUFFLED, make_tracker
from reporting import derive, load, metrics, render
from reporting.load import Job, Settings

TODAY = date(2026, 10, 28)  # a Wednesday: this week is Oct 26-28, last week Oct 19-25


def _fingerprint(path):
    return path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest()


def plant(path, rows):
    """Write {header: value} rows over the fake tracker, starting at row 2 (by header name)."""
    wb = load_workbook(path)
    ws = wb["Sheet1"]
    col = {c.value: c.column for c in ws[1]}
    for r, values in enumerate(rows, start=2):
        for header, value in values.items():
            cell = ws.cell(r, col[header])
            cell.value = value  # ws.cell(r, c, None) would leave the old value
            if header == "URL":  # a plain link is text + hyperlink like real rows; a formula or blank has none
                cell.hyperlink = value if value and not str(value).startswith("=") else None
            if isinstance(value, datetime):
                cell.number_format = "M/d/yyyy"
    wb.save(path)


# ---------- load ----------

def test_excel_rows_by_header_name(tracker_file):
    plant(tracker_file, [{"Company": "Acme", "Position": "Backend Engineer", "URL": "https://jobs.lever.co/acme/1",
                          "Posted date/ Source": "Posted today-Jobright", "Date": datetime(2026, 10, 27),
                          "Outcome": "In progress", "Stage": "Interview", "Last Updated": datetime(2026, 10, 27),
                          "Remarks": "referred by Sam"}])
    jobs = load.excel_rows(tracker_file)
    assert len(jobs) == 896
    assert jobs[0] == Job(row=2, company="Acme", position="Backend Engineer", url="https://jobs.lever.co/acme/1",
                          posted="Posted today-Jobright", applied=date(2026, 10, 27), outcome="In progress",
                          stage="Interview", last_updated=date(2026, 10, 27), remarks="referred by Sam",
                          url_is_formula=False)
    assert jobs[1].company == "Company 3" and jobs[1].last_updated is None and jobs[1].stage == ""


def test_excel_rows_in_a_shuffled_layout(tmp_path):
    path = make_tracker(tmp_path / "shuffled.xlsx", headers=SHUFFLED)
    plant(path, [{"Company": "Acme", "Position": "SRE", "Outcome": "Rejected", "Stage": "Online Assessment",
                  "Remarks": "no visa"}])
    job = load.excel_rows(path)[0]
    assert (job.company, job.position, job.outcome, job.stage, job.remarks) == (
        "Acme", "SRE", "Rejected", "Online Assessment", "no visa")
    assert job.url == "https://example.com/jobs/2"


def test_hyperlink_formula_gives_the_target(tracker_file):
    plant(tracker_file, [{"URL": '=HYPERLINK("https://jobs.ashbyhq.com/acme/9", "click here")'},
                         {"URL": None, "Company": "No Link Co"}])
    first, second = load.excel_rows(tracker_file)[:2]
    assert first.url == "https://jobs.ashbyhq.com/acme/9" and first.url_is_formula
    assert second.url == "" and not second.url_is_formula


def test_hyperlink_target_wins_over_link_text(tracker_file):
    wb = load_workbook(tracker_file)
    ws = wb["Sheet1"]
    ws.cell(2, 3, "click here").hyperlink = "https://jobs.lever.co/acme/2"
    wb.save(tracker_file)
    assert load.excel_rows(tracker_file)[0].url == "https://jobs.lever.co/acme/2"


def test_reading_never_changes_the_tracker(tracker_file):
    before = _fingerprint(tracker_file)
    load.excel_rows(tracker_file)
    assert _fingerprint(tracker_file) == before


def test_sheet_rows_use_the_formula_target():
    ws = FakeWorksheet([
        HEADER,
        ["Acme", "Data Engineer", "click here", "Posted today-Tsenta", "10/27/2026", "Awaiting reply", "Interview",
         "10/27/2026", "", "", ""],
        ["Zenith", "SWE", "https://jobs.lever.co/zenith/1", "LinkedIn", "10/1/2026", "Applied", "", "", "", "", ""],
    ], formulas={2: '=HYPERLINK("https://jobs.lever.co/acme/3", "click here")'})
    acme, zenith = load.sheet_rows(ws)
    assert acme.url == "https://jobs.lever.co/acme/3" and not acme.url_is_formula  # formulas: Excel check only
    assert acme.applied == date(2026, 10, 27) and acme.last_updated == date(2026, 10, 27)
    assert zenith.url == "https://jobs.lever.co/zenith/1" and zenith.last_updated is None
    assert ws.updates == []


def test_settings_defaults_when_no_file(tmp_path):
    s = load.load_settings(tmp_path / "missing.yaml")
    assert s == Settings()
    assert (s.weekly_target, s.follow_up_days, s.no_reply_days, s.many_applications, s.min_group,
            s.weeks_shown) == (25, 14, 30, 5, 20, 12)
    assert s.sources == ["Jobright", "Tsenta", "LinkedIn"]
    assert list(s.title_types) == ["Backend", "DevOps/Cloud/SRE/Platform", "AI/ML/Data", "Software (general)"]
    assert list(s.remark_tags) == ["No sponsorship", "Clearance/citizenship", "Referral"]


def test_settings_file_overrides_only_its_keys(tmp_path):
    path = tmp_path / "report.yaml"
    path.write_text("weekly_target: 10\nsources: [Handshake]\n", encoding="utf-8")
    s = load.load_settings(path)
    assert s.weekly_target == 10 and s.sources == ["Handshake"] and s.follow_up_days == 14


def test_settings_unknown_key_is_named(tmp_path):
    path = tmp_path / "report.yaml"
    path.write_text("weekly_targe: 10\n", encoding="utf-8")
    with pytest.raises(SystemExit, match="weekly_targe"):
        load.load_settings(path)


def test_repo_report_yaml_matches_the_defaults():
    from pathlib import Path
    assert load.load_settings(Path(__file__).resolve().parent.parent / "report.yaml") == Settings()


# ---------- derive ----------

def job(row=2, company="Acme", position="Software Engineer", url="https://jobs.lever.co/acme/1",
        posted="Posted today-Jobright", applied=date(2026, 10, 1), outcome="Applied", stage="",
        last_updated=None, remarks="", url_is_formula=False) -> Job:
    return Job(row, company, position, url, posted, applied, outcome, stage, last_updated, remarks, url_is_formula)


S = Settings()


@pytest.mark.parametrize("posted, source", [
    ("Posted today-Jobright", "Jobright"),
    ("Sameday- Tsenta", "Tsenta"),
    ("Posted October 8, 2026-linkedin", "LinkedIn"),
    ("Posted 2 days ago-Jobright via LinkedIn", "Jobright"),  # first of `sources`, not first in the text
    ("Posted today-Other", "Other"),
    ("Posted today", "Other"),
    ("", "Other"),
    ("Posted today-LinkedInish", "Other"),  # whole word only
])
def test_source(posted, source):
    assert derive.source(posted, S.sources) == source


@pytest.mark.parametrize("posted, named", [
    ("Posted today-Tsenta", True), ("Posted today-Other", True), ("Posted today", False), ("", False),
])
def test_source_named(posted, named):
    assert derive.source_named(posted, S.sources) is named


@pytest.mark.parametrize("url, system", [
    ("https://job-boards.greenhouse.io/acme/jobs/1", "Greenhouse"),
    ("https://jobs.ashbyhq.com/acme/1", "Ashby"),
    ("https://acme.wd5.myworkdayjobs.com/External/job/X_R1", "Workday"),
    ("https://careers-acme.icims.com/jobs/1/job", "iCIMS"),
    ("https://example.com/jobs/1", "Other"),
    ("", "Unknown"),
])
def test_job_system(url, system):
    assert derive.job_system(url) == system


@pytest.mark.parametrize("title, level", [
    ("Software Engineer I", "Early career"),
    ("Software Engineer 1", "Early career"),
    ("New Grad Software Engineer", "Early career"),
    ("Software Engineer, Early Career", "Early career"),
    ("Jr. Developer", "Early career"),
    ("Software Engineering Intern", "Early career"),
    ("Software Engineer II", "Other"),
    ("I/O Engineer", "Other"),  # " I" only as the last word
    ("Senior Software Engineer", "Other"),
    ("Internal Tools Engineer", "Other"),  # whole word: "intern" is not "internal"
    ("", "Other"),
])
def test_title_level(title, level):
    assert derive.title_level(title, S.early_career_words) == level


@pytest.mark.parametrize("title, kind", [
    ("Backend Software Engineer", "Backend"),  # Backend comes before Software (general)
    ("Software Engineer, API Platform", "Backend"),
    ("Site Reliability Engineer", "DevOps/Cloud/SRE/Platform"),
    ("Cloud Engineer I", "DevOps/Cloud/SRE/Platform"),
    ("AI Engineer", "AI/ML/Data"),
    ("Data Engineer", "AI/ML/Data"),
    ("SDE 1", "Software (general)"),
    ("Product Manager", "Other"),
    ("Email Marketing", "Other"),  # "ai" and "ml" as whole words only
])
def test_title_type(title, kind):
    assert derive.title_type(title, S.title_types) == kind


@pytest.mark.parametrize("outcome, stage, group", [
    ("In progress", "Online Assessment", "In progress"),
    ("Awaiting reply", "Interview", "Awaiting reply"),
    ("Rejected", "Initial Screening", "Rejected after response"),
    ("Rejected", "", "Rejected"),
    ("Applied", "", "No news"),
    ("", "", "No news"),
    ("Applied", "Interview", "Awaiting reply"),  # the old model: a Stage meant waiting on them
])
def test_status_group(outcome, stage, group):
    assert derive.status_group(outcome, stage) == group


def test_remark_tags():
    assert derive.remark_tags("10/02/2026 no visa sponsorship | referred by Sam", S.remark_tags) == [
        "No sponsorship", "Referral"]
    assert derive.remark_tags("US Citizens only", S.remark_tags) == ["Clearance/citizenship"]
    assert derive.remark_tags("", S.remark_tags) == []


def test_facts_of_a_row():
    j = job(company="Acme, Inc.", position="Backend Engineer I", posted="Posted today-Tsenta",
            applied=date(2026, 10, 8), outcome="Awaiting reply", stage="Interview",
            last_updated=date(2026, 10, 14), remarks="referral from Sam")
    f = derive.facts(j, S, TODAY)
    assert f.job is j
    assert (f.source, f.system, f.level, f.kind, f.responded, f.group) == (
        "Tsenta", "Lever", "Early career", "Backend", True, "Awaiting reply")
    assert (f.age, f.since_update, f.tags, f.company_key) == (20, 14, ["Referral"], "acme")


def test_facts_with_unknown_dates():
    f = derive.facts(job(applied=None, last_updated=None), S, TODAY)
    assert f.age is None and f.since_update is None and not f.responded and f.group == "No news"


# ---------- metrics ----------

def F(**kw) -> derive.Facts:
    return derive.facts(job(**kw), S, TODAY)


def d(month, day, year=2026):
    return date(year, month, day)


def test_change_text():
    assert metrics.change_text(19, 6) == "↑ 13 from last week"
    assert metrics.change_text(2, 5) == "↓ 3 from last week"
    assert metrics.change_text(4, 4) == "same as last week"


@pytest.mark.parametrize("r, n, text", [
    (5, 128, "5 of 128 (3.9%)"), (1, 20, "1 of 20 (5%)"), (0, 25, "0 of 25 (0%)"), (3, 12, "3 of 12"),
])
def test_rate_text(r, n, text):
    assert metrics.rate_text(r, n, 20) == text


WEEK_ROWS = [
    F(company="A", applied=d(10, 26)), F(company="B", applied=d(10, 28)), F(company="C", applied=d(10, 28)),
    F(company="D", applied=d(10, 19)), F(company="E", applied=d(10, 25)),  # last week
    F(company="F", applied=d(10, 18)), F(company="G", applied=None),
    F(company="Delta", applied=d(10, 1), outcome="Awaiting reply", stage="Interview", last_updated=d(10, 27)),
    F(company="Zeta", applied=d(10, 2), outcome="In progress", stage="Online Assessment", last_updated=d(10, 25)),
    F(company="Eta", applied=d(9, 2), outcome="Awaiting reply", stage="Initial Screening"),  # no Last Updated
]


def test_this_week():
    week = metrics.this_week(WEEK_ROWS, TODAY)
    assert (week.applications, week.last_week, week.in_progress, week.awaiting) == (3, 2, 1, 2)
    assert [f.job.company for f in week.new_responses] == ["Delta"]


def test_summary_sentences():
    assert metrics.summary(WEEK_ROWS, S, TODAY) == [
        "This week: 3 applications (target 25, last week 2). 1 new response: Delta (Interview).",
        "1 in progress and 2 awaiting reply.",
    ]


def test_summary_leaves_out_empty_sentences():
    assert metrics.summary([], S, TODAY) == ["This week: 0 applications (target 25, last week 0). No new responses."]
    one = [F(applied=d(10, 27))]
    assert metrics.summary(one, S, TODAY)[0] == (
        "This week: 1 application (target 25, last week 0). No new responses.")


def _source_rows(source, n, responded, applied=d(1, 5)):
    return [F(company=f"{source} {i}", posted=f"Posted today-{source}", applied=applied,
              outcome="Rejected", stage="Interview" if i < responded else "") for i in range(n)]


def test_summary_best_source():
    rows = _source_rows("Tsenta", 20, 2) + _source_rows("Jobright", 25, 2) + _source_rows("LinkedIn", 5, 3)
    assert metrics.summary(rows, S, TODAY) == [
        "This week: 0 applications (target 25, last week 0). No new responses.",
        "Best source so far: Tsenta, 2 responses from 20 applications.",  # 10% beats 8%; LinkedIn too few
    ]


def test_summary_no_best_source_with_one_response():
    rows = _source_rows("Tsenta", 30, 1)
    assert len(metrics.summary(rows, S, TODAY)) == 1


def test_summary_two_new_responses():
    rows = [F(company="Acme", stage="Interview", outcome="In progress", last_updated=d(10, 26)),
            F(company="Zen", stage="Online Assessment", outcome="Awaiting reply", last_updated=d(10, 28))]
    assert metrics.summary(rows, S, TODAY)[0].endswith(
        "2 new responses: Acme (Interview), Zen (Online Assessment).")


ATTENTION_ROWS = [
    F(company="Beta", outcome="In progress", stage="Online Assessment", last_updated=d(10, 25)),
    F(company="Gamma", outcome="In progress", stage="Interview", last_updated=d(10, 18)),
    F(company="Nu", outcome="In progress", stage="Interview"),
    F(company="Omicron", outcome="Awaiting reply", stage="Initial Screening", last_updated=d(10, 10)),
    F(company="Delta", outcome="Awaiting reply", stage="Interview", last_updated=d(10, 1)),
    F(company="Zeta", outcome="Awaiting reply", stage="Online Assessment", last_updated=d(10, 20)),  # 8 days: fine
    F(company="Epsilon", outcome="Awaiting reply", stage="Initial Screening"),
    *[F(company="Acme Inc" if i < 3 else "Acme", applied=d(9, 1 + i)) for i in range(5)],
    *[F(company="Pi", applied=d(9, 1 + i)) for i in range(4)],
    F(company="Pi", outcome="Rejected", stage="Online Assessment"),  # one response: not flagged
    *[F(company="Rho", applied=d(9, 1 + i)) for i in range(4)],  # only 4
]


def test_attention_order():
    assert metrics.attention(ATTENTION_ROWS, S, TODAY) == [
        "Gamma: Interview to do since 10 days",
        "Beta: Online Assessment to do since 3 days",
        "Nu: Interview to do",
        "Delta: Interview, no reply for 27 days",
        "Omicron: Initial Screening, no reply for 18 days",
        "Epsilon: Initial Screening, waiting (date unknown)",
        "Acme: 5 applications, no response",
    ]


def test_attention_cap():
    lines = metrics.attention(ATTENTION_ROWS, S, TODAY)
    shown, more = metrics.cap(lines)
    assert shown == lines[:5] and more == lines[5:]
    assert metrics.cap(["a"]) == (["a"], [])
    assert metrics.attention([F()], S, TODAY) == []


def test_weekly_series():
    rows = [F(applied=d(10, 26), stage="Interview", outcome="Awaiting reply"), F(applied=d(10, 28)),
            F(applied=d(10, 19)), F(applied=d(8, 10)), F(applied=d(8, 9)), F(applied=None)]
    weeks = metrics.weekly(rows, S, TODAY)
    assert len(weeks) == 12
    assert weeks[0] == metrics.Week(d(8, 10), 1, 0)
    assert weeks[-2] == metrics.Week(d(10, 19), 1, 0)
    assert weeks[-1] == metrics.Week(d(10, 26), 2, 1)
    assert sum(w.applications for w in weeks) == 4


def test_monthly_table():
    rows = ([F(applied=d(9, 1 + i)) for i in range(17)]
            + [F(applied=d(9, 20), outcome="Rejected", stage="Interview"),
               F(applied=d(9, 21), outcome="Rejected"), F(applied=d(9, 22), outcome="Rejected")]
            + [F(applied=d(10, 2)), F(applied=d(10, 3), outcome="In progress", stage="Interview")]
            + [F(applied=d(7, 4)), F(applied=None)])
    assert metrics.monthly(rows, S) == [
        metrics.Month(d(10, 1), 2, 1, 0, "1 of 2"),
        metrics.Month(d(9, 1), 20, 1, 3, "1 of 20 (5%)"),
        metrics.Month(d(7, 1), 1, 0, 0, "0 of 1"),
    ]


def test_responses_newest_first():
    rows = [F(company="Old", applied=d(8, 1), stage="Interview", outcome="Rejected"),
            F(company="None", applied=None, stage="Interview", outcome="Awaiting reply"),
            F(company="New", applied=d(10, 1), stage="Online Assessment", outcome="In progress"),
            F(company="No stage", applied=d(10, 20))]
    assert [f.job.company for f in metrics.responses(rows)] == ["New", "Old", "None"]


def test_top_companies():
    rows = ([F(company="Acme Inc" if i % 2 else "Acme", applied=d(10, 16 + i)) for i in range(6)]  # last Oct 21
            + [F(company="Beta", applied=d(10, 18)), F(company="Beta", applied=d(10, 19)),
               F(company="Beta", applied=d(10, 20), outcome="In progress", stage="Online Assessment")]
            + [F(company="Gamma", applied=d(10, 20)), F(company="Gamma", applied=d(10, 21), outcome="Rejected"),
               F(company="Gamma", applied=d(10, 22), outcome="Rejected", stage="Interview")]
            + [F(company=f"S{i}", applied=d(10, i)) for i in range(1, 10)]
            + [F(company="", applied=d(10, 27))])
    top = metrics.top_companies(rows, S)
    assert [c.name for c in top] == ["Acme Inc", "Gamma", "Beta", "S9", "S8", "S7", "S6", "S5", "S4", "S3"]
    assert top[0] == metrics.CompanyRow("Acme Inc", 6, 0, 0, 0, d(10, 21), "muted")  # name of the newest row
    assert top[1] == metrics.CompanyRow("Gamma", 3, 1, 0, 2, d(10, 22), "green")
    assert top[2] == metrics.CompanyRow("Beta", 3, 1, 1, 0, d(10, 20), "green")
    assert top[3] == metrics.CompanyRow("S9", 1, 0, 0, 0, d(10, 9), "")


def test_no_reply():
    rows = [F(company="A", applied=d(9, 27)), F(company="B", applied=d(9, 28)),  # 31 days / exactly 30
            F(company="C", applied=d(8, 5)), F(company="D", applied=d(8, 20)),
            F(company="E", applied=d(8, 1), outcome="Rejected"),
            F(company="F", applied=d(7, 1), outcome="Awaiting reply", stage="Interview"),
            F(company="G", applied=d(7, 1), stage="Interview"), F(company="H", applied=None)]
    total, months = metrics.no_reply(rows, S, TODAY)
    assert total == 3
    assert [(m, [f.job.company for f in fs]) for m, fs in months] == [(d(9, 1), ["A"]), (d(8, 1), ["D", "C"])]


def test_details_tables():
    rows = ([F(posted="Posted today-Jobright", outcome="Rejected", stage="Interview" if i < 2 else "")
             for i in range(3)]
            + [F(posted="Posted today-Jobright") for _ in range(17)]
            + [F(posted="Posted today-Tsenta", remarks="no visa"), F(posted="Posted today-Tsenta")])
    tables = dict(metrics.details(rows, S))
    assert list(tables) == ["Source", "Job system", "Title level", "Title type", "Remark tags"]
    assert tables["Source"] == [metrics.GroupRow("Jobright", 20, 2, "10%", 3),
                                metrics.GroupRow("Tsenta", 2, 0, "too few to compare", 0)]
    assert tables["Remark tags"] == [metrics.GroupRow("No sponsorship", 1, 0, "too few to compare", 0)]
    assert tables["Job system"] == [metrics.GroupRow("Lever", 22, 2, "9.1%", 3)]


def test_data_check():
    rows = [F(row=2, company="Ok"),
            F(row=3, company="NoPos", position=""),
            F(row=4, company="NoUrl", url=""),
            F(row=5, company="NoSrc", posted="Posted today"),
            F(row=6, company="Formula", url_is_formula=True),
            F(row=7, company="Old", outcome="Applied", stage="Interview"),
            F(row=8, company="Stuck", outcome="In progress"),
            F(row=9, company="Other", posted="Posted today-Other")]
    issues = metrics.data_check(rows, S, excel=True)
    assert [(name, [(f.job.row, f.job.company) for f in fs]) for name, fs in issues] == [
        ("Blank Position", [(3, "NoPos")]),
        ("Blank URL", [(4, "NoUrl")]),
        ("No source found in Posted date/ Source", [(5, "NoSrc")]),
        ("URL stored as a formula", [(6, "Formula")]),
        ("Stage set but Outcome Applied", [(7, "Old")]),
        ("In progress with no Stage", [(8, "Stuck")]),
    ]
    assert "URL stored as a formula" not in dict(metrics.data_check(rows, S, excel=False))
    assert metrics.data_check([F()], S, excel=True) == []


# ---------- render ----------

HEADINGS = ["Summary", "This week", "Needs your attention", "Weekly", "Monthly", "Your responses (",
            "Top 10 companies, with outcomes", "No reply after 30 days", "More details", "Data check ("]
GENERATED = datetime(2026, 10, 28, 9, 5)


def page(rows, excel=True):
    return render.page(rows, S, TODAY, excel=excel, source_name="tracker.xlsx", generated=GENERATED)


def test_sections_in_order():
    html = page(ATTENTION_ROWS + WEEK_ROWS)
    headings = re.findall(r"<h2[^>]*>(.*?)</h2>", html)
    assert len(headings) == len(HEADINGS)
    for heading, start in zip(headings, HEADINGS):
        assert heading.startswith(start), (heading, start)
    assert "Your responses (11)" in headings[5]


def test_tracker_text_is_escaped():
    html = page([F(company="<script>alert(1)</script>", position="A & B", stage="Interview", outcome="In progress",
                   remarks="<b>x</b>")])
    assert "<script>alert" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "A &amp; B" in html


def test_no_external_requests():
    html = page(ATTENTION_ROWS + WEEK_ROWS + [F(company="Bad", url="javascript:alert(1)", stage="Interview")])
    for banned in ("<script", "<link", "@import", "url(", "javascript:"):
        assert banned not in html.lower(), banned
    links = re.findall(r'<a href="([^"]*)"[^>]*>', html)
    assert links and all(u.startswith("https://") for u in links)
    assert all('target="_blank"' in a and 'rel="noopener noreferrer"' in a for a in re.findall(r"<a [^>]*>", html))
    without_links = re.sub(r'<a href="[^"]*"', "", html)
    assert "http" not in without_links


def test_job_link_shows_the_domain():
    html = page([F(url="https://www.jobs.lever.co/acme/1", stage="Interview", outcome="Awaiting reply")])
    assert re.search(r'<a href="https://www.jobs.lever.co/acme/1"[^>]*>jobs.lever.co</a>', html)


def test_attention_card_cap_and_empty():
    html = page(ATTENTION_ROWS)
    assert "Acme: 5 applications, no response" in html and "+2 more" in html
    assert "Nothing needs your attention." in page([F()])


def test_this_week_card_and_summary():
    html = page(WEEK_ROWS)
    assert "This week: 3 applications (target 25, last week 2). 1 new response: Delta (Interview)." in html
    assert "3 · ↑ 1 from last week" in html


def test_weekly_chart():
    html = page(WEEK_ROWS)
    assert html.count("<rect") >= 12
    assert "<title>Week of Oct 26: 3 applications, 0 responses</title>" in html
    assert 'stroke-dasharray' in html  # the target line
    assert "Oct 26" in html


def test_footer_and_empty_tracker():
    html = page([])
    assert "Generated 10/28/2026 9:05 AM from tracker.xlsx · 0 rows" in html
    assert "Data check (0)" in html

