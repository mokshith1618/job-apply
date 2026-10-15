import faulthandler
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

import tracker
from conftest import KNOWN_POSITION, KNOWN_ROW, KNOWN_URL, KNOWN_COMPANY
from tracker import TrackerRow

TODAY = date(2026, 10, 13)
EXPECTED_ROW = 898  # first row after the last Company value (fake_tracker.LAST_ROW + 1, as in SPEC.md)
HEADERS = ["Company", "Position", "URL", "Posted date/ Source", "Date", "Outcome",
           "Remarks", "Experience", "Key Skills"]
ASHBY_ROW = KNOWN_ROW
ASHBY_URL = KNOWN_URL + "?utm_source=LinkedInPaid"


def new_row(**kw) -> TrackerRow:
    base = dict(company="Acme Robotics", position="Software Engineer I",
                url="https://job-boards.greenhouse.io/acme/jobs/1234567",
                posted="Posted today-Tsenta", applied_on=TODAY)
    return TrackerRow(**{**base, **kw})


def company_rows(path) -> int:
    ws = load_workbook(path)["Sheet1"]
    return sum(1 for (v,) in ws.iter_rows(min_row=2, max_col=1, values_only=True) if v not in (None, ""))


def dv_covering(ws, coord):
    return [dv for dv in ws.data_validations.dataValidation if coord in dv.sqref]


# ---------- Done when: a new row appears at row 898 with the right formats ----------

def test_new_row_lands_at_898_with_right_values_and_formats(tracker_file):
    before = load_workbook(tracker_file)["Sheet1"]
    row_897_before = [c.value for c in before[897]]
    count_before = company_rows(tracker_file)

    r = tracker.append_row(tracker_file, new_row(), backup_dir=tracker_file.parent / "backups")
    assert r == EXPECTED_ROW

    ws = load_workbook(tracker_file)["Sheet1"]
    assert ws["A898"].value == "Acme Robotics"
    assert ws["B898"].value == "Software Engineer I"
    # URL as text and as a hyperlink
    assert ws["C898"].value == "https://job-boards.greenhouse.io/acme/jobs/1234567"
    assert ws["C898"].hyperlink is not None
    assert ws["C898"].hyperlink.target == "https://job-boards.greenhouse.io/acme/jobs/1234567"
    assert ws["C898"].font.color.rgb == ws["C897"].font.color.rgb  # same link look as row above
    assert ws["C898"].font.u == "single"
    assert ws["D898"].value == "Posted today-Tsenta"
    # real date, M/d/yyyy
    assert isinstance(ws["E898"].value, datetime)
    assert ws["E898"].value.date() == TODAY
    assert ws["E898"].number_format == "M/d/yyyy"
    # Outcome is a dropdown value and the cell is covered by the dropdown
    assert ws["F898"].value == "Applied"
    f_dv = dv_covering(ws, "F898")
    assert f_dv and f_dv[0].type == "list"
    assert "Applied" in f_dv[0].formula1.split(",")[0]
    assert dv_covering(ws, "E898")  # date check still applies
    # G, H, I empty
    assert [ws.cell(898, c).value for c in (7, 8, 9)] == [None, None, None]
    # nothing else moved
    assert [c.value for c in ws[897]] == row_897_before
    assert company_rows(tracker_file) == count_before + 1
    assert ws["A899"].value is None


def test_second_row_goes_to_899(tracker_file):
    tracker.append_row(tracker_file, new_row())
    assert tracker.append_row(tracker_file, new_row(company="Second Co")) == EXPECTED_ROW + 1


# ---------- Done when: the file still opens in Excel ----------

def test_saved_file_is_structurally_intact(tracker_file):
    wb0 = load_workbook(tracker_file)
    style_before = wb0["Sheet1"].tables["Table1"].tableStyleInfo.name
    links_before = sum(1 for row in wb0["Sheet1"].iter_rows() for c in row if c.hyperlink)

    tracker.append_row(tracker_file, new_row())

    with zipfile.ZipFile(tracker_file) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert "xl/workbook.xml" in names and any(n.startswith("xl/tables/") for n in names)
    wb = load_workbook(tracker_file)
    assert wb.sheetnames == ["Sheet1"]
    ws = wb["Sheet1"]
    table = ws.tables["Table1"]
    assert table.ref == "A1:I1000"
    assert [col.name for col in table.tableColumns] == HEADERS
    assert [c.value for c in ws[1]][:9] == HEADERS
    assert table.tableStyleInfo.name == style_before
    assert {dv.type for dv in ws.data_validations.dataValidation} == {"list", "custom"}
    assert ws.freeze_panes == "A2"
    links_after = sum(1 for row in ws.iter_rows() for c in row if c.hyperlink)
    assert links_after == links_before + 1


def test_opens_in_real_excel(tracker_file):
    """Open the written file in a hidden Excel instance (skipped if Excel/pywin32 is missing)."""
    win32 = pytest.importorskip("win32com.client")
    tracker.append_row(tracker_file, new_row())
    try:
        excel = win32.DispatchEx("Excel.Application")
    except Exception as exc:  # noqa: BLE001
        pytest.skip(f"Excel not available: {exc}")
    excel.Visible = False
    excel.DisplayAlerts = False
    # Excel raises handled COM exceptions (0x80010108) on shutdown; faulthandler would dump them as noise.
    faulthandler_was_on = faulthandler.is_enabled()
    faulthandler.disable()
    try:
        wb = excel.Workbooks.Open(str(tracker_file), UpdateLinks=0, ReadOnly=True)
        ws = wb.Worksheets("Sheet1")
        seen = {
            "company": ws.Range("A898").Value,
            "links": ws.Range("C898").Hyperlinks.Count,
            "date_format": ws.Range("E898").NumberFormat,
            "date": ws.Range("E898").Value.date(),
            "outcome": ws.Range("F898").Value,
            "outcome_validation": ws.Range("F898").Validation.Type,
            "tables": ws.ListObjects.Count,
            "table_range": ws.ListObjects("Table1").Range.Address,
        }
        wb.Close(SaveChanges=False)
        del ws, wb  # release COM objects before Excel quits
    finally:
        excel.Quit()
        del excel
        if faulthandler_was_on:
            faulthandler.enable()

    assert seen["company"] == "Acme Robotics"
    assert seen["links"] == 1
    # Excel shows the format with this PC's date separator ("/" in M/d/yyyy means "locale separator")
    assert seen["date_format"].replace("-", "/").replace(".", "/") == "m/d/yyyy"
    assert seen["date"] == TODAY
    assert seen["outcome"] == "Applied"
    assert seen["outcome_validation"] == 3  # xlValidateList: Outcome dropdown
    assert seen["tables"] == 1
    assert seen["table_range"] == "$A$1:$I$1000"


# ---------- Done when: the duplicate check catches a URL already in the tracker ----------

def test_duplicate_exact_url(dup_tracker):
    hits = tracker.find_duplicates(dup_tracker, url=ASHBY_URL)
    assert [r for r, _ in hits] == [ASHBY_ROW]
    assert hits[0][1][0] == KNOWN_COMPANY


@pytest.mark.parametrize("variant", [
    KNOWN_URL,
    KNOWN_URL + "?utm_source=Tsenta&utm_medium=x",
    KNOWN_URL.replace("https:", "http:") + "/application",
    KNOWN_URL + "/?source=LinkedIn#top",
])
def test_duplicate_url_ignores_tracking_params(dup_tracker, variant):
    assert [r for r, _ in tracker.find_duplicates(dup_tracker, url=variant)] == [ASHBY_ROW]


def test_duplicate_company_and_position(dup_tracker):
    hits = tracker.find_duplicates(dup_tracker, url="https://example.com/other-link",
                                   company="  example LABS, inc. ", position=KNOWN_POSITION)
    assert ASHBY_ROW in [r for r, _ in hits]


def test_no_duplicate_for_new_job(tracker_file):
    assert tracker.find_duplicates(tracker_file, url="https://job-boards.greenhouse.io/acme/jobs/1234567",
                                   company="Acme Robotics", position="Software Engineer I") == []


def test_duplicate_found_after_append(tracker_file):
    tracker.append_row(tracker_file, new_row())
    hits = tracker.find_duplicates(tracker_file, url="https://job-boards.greenhouse.io/acme/jobs/1234567?gh_src=abc")
    assert [r for r, _ in hits] == [EXPECTED_ROW]


# ---------- supporting behaviour ----------

def test_normalize_url():
    n = tracker.normalize_url
    assert n("https://www.Example.com/jobs/1/?utm_source=x&b=2&a=1#frag") == "example.com/jobs/1?a=1&b=2"
    assert n("https://boards.greenhouse.io/x/jobs/1?gh_src=abc") == "boards.greenhouse.io/x/jobs/1"
    assert n("https://jobs.lever.co/acme/abc/apply?lever-source=LinkedIn") == "jobs.lever.co/acme/abc"
    assert n("https://site.com/careers?gh_jid=55") == "site.com/careers?gh_jid=55"
    assert n("https://careers-x.icims.com/jobs/1/title/job?mobile=false&width=1380&height=500&bga=true"
             "&needsRedirect=false&jan1offset=-420&jun1offset=-420") == "careers-x.icims.com/jobs/1/title/job"


def test_format_posted():
    assert tracker.format_posted(TODAY, "Tsenta", TODAY) == "Posted today-Tsenta"
    assert tracker.format_posted(date(2026, 10, 8), "LinkedIn", TODAY) == "Posted October 8, 2026-LinkedIn"
    assert tracker.format_posted(None, "Other", TODAY) == "Other"


def test_backup_is_made_before_write(tracker_file):
    original = tracker_file.read_bytes()
    backups = tracker_file.parent / "backups"
    tracker.append_row(tracker_file, new_row(), backup_dir=backups)
    files = list(backups.glob("tracker-*.xlsx"))
    assert len(files) == 1
    assert files[0].name.startswith("tracker-") and len(files[0].stem) == len("tracker-YYYYMMDD-HHMMSS")
    assert files[0].read_bytes() == original  # backup is the pre-write file


def test_default_backup_dir_is_next_to_tracker(tracker_file):
    tracker.append_row(tracker_file, new_row())
    assert len(list((tracker_file.parent / "backups").glob("tracker-*.xlsx"))) == 1


def test_table_and_validations_extend_when_full(tracker_file):
    """Shrink Table1 and the E/F validations to end at row 897, then append: they must grow."""
    wb = load_workbook(tracker_file)
    ws = wb["Sheet1"]
    ws.tables["Table1"].ref = "A1:I897"
    for dv in ws.data_validations.dataValidation:
        col = str(dv.sqref)[0]
        dv.sqref = tracker.MultiCellRange(f"{col}2:{col}897")
    wb.save(tracker_file)

    r = tracker.append_row(tracker_file, new_row())
    assert r == EXPECTED_ROW
    ws = load_workbook(tracker_file)["Sheet1"]
    assert ws.tables["Table1"].ref == f"A1:I{EXPECTED_ROW + tracker.TABLE_GROWTH}"
    assert dv_covering(ws, "F898")[0].type == "list"
    assert dv_covering(ws, "E898")[0].type == "custom"
    assert dv_covering(ws, f"F{EXPECTED_ROW + tracker.TABLE_GROWTH}")


def test_locked_tracker_keeps_row_pending_then_flushes(tracker_file, monkeypatch):
    pending = tracker.pending_path_for(tracker_file)
    tracker.save_pending(pending, new_row())

    count_before = company_rows(tracker_file)
    real_replace = tracker.os.replace

    def locked(*_args, **_kwargs):
        raise PermissionError(13, "The process cannot access the file")

    monkeypatch.setattr(tracker.os, "replace", locked)
    with pytest.raises(tracker.TrackerLocked):
        tracker.flush_pending(tracker_file)
    assert [r.company for r in tracker.load_pending(pending)] == ["Acme Robotics"]
    assert company_rows(tracker_file) == count_before  # unchanged
    assert not list(tracker_file.parent.glob("~*.tmp.xlsx"))  # temp file cleaned up

    monkeypatch.setattr(tracker.os, "replace", real_replace)
    written = tracker.flush_pending(tracker_file)
    assert [r for r, _ in written] == [EXPECTED_ROW]
    assert not pending.exists()
    assert load_workbook(tracker_file)["Sheet1"]["A898"].value == "Acme Robotics"


def test_pending_row_round_trip(tmp_path):
    p = tmp_path / "pending_rows.json"
    tracker.save_pending(p, new_row())
    tracker.save_pending(p, new_row(company="B"))
    rows = tracker.load_pending(p)
    assert [r.company for r in rows] == ["Acme Robotics", "B"]
    assert rows[0].applied_on == TODAY
