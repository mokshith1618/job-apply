"""The tracker in Google Sheets: duplicate check and row writing, same rules as tracker.py.

Access is through a Google Cloud service account (google-key.json); the Sheet is shared with
its email as Editor. Every function takes a gspread Worksheet, so tests can pass a fake one.
Columns are found by header name (tracker.column_map). Text is written RAW, so it is stored
exactly as given; only dates and the =HYPERLINK URL are entered as if typed (USER_ENTERED).
"""
from __future__ import annotations

import csv
import functools
import re
from datetime import datetime
from pathlib import Path

from openpyxl.utils.cell import get_column_letter

import tracker
from tracker import TrackerRow

DATE_PATTERN = tracker.DATE_FORMAT  # M/d/yyyy, same as the Excel file
SHEET_GROWTH = 100  # rows added when the sheet has no empty row left
HYPERLINK_TARGET = re.compile(r'^=HYPERLINK\(\s*"((?:[^"]|"")*)"', re.IGNORECASE)
ENTERED = ("URL", *tracker.DATE_FIELDS)  # written USER_ENTERED; every other field RAW


class SheetUnavailable(Exception):
    """The Sheet can't be reached: offline, Google error, bad key, or not shared with the service account."""


class SheetMismatch(tracker.TrackerLayoutError):
    """Row 1 lacks a tracker header, so this is probably the wrong sheet or tab. Nothing is written."""


def _network_errors() -> tuple:
    import gspread
    import requests
    from google.auth.exceptions import GoogleAuthError
    return gspread.exceptions.GSpreadException, GoogleAuthError, requests.exceptions.RequestException


def _guard(fn):
    """Turn any Google / network error into SheetUnavailable, so callers only handle one kind."""
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        try:
            return fn(*args, **kwargs)
        except _network_errors() as exc:
            raise SheetUnavailable(f"{type(exc).__name__}: {exc}") from exc
    return wrapper


@_guard
def open_sheet(sheet_id: str, tab: str, key_file):
    import gspread
    key_file = Path(key_file)
    if not key_file.exists():
        raise SheetUnavailable(f"service account key not found: {key_file}")
    return gspread.service_account(filename=str(key_file)).open_by_key(sheet_id).worksheet(tab)


def columns(values: list[list]) -> dict[str, int]:
    """Header name -> column number from row 1. A missing header raises SheetMismatch naming it."""
    try:
        return tracker.column_map(values[0] if values else [])
    except tracker.TrackerLayoutError as exc:
        raise SheetMismatch(str(exc)) from exc


def _pad(row: list, n: int) -> list:
    return list(row[:n]) + [""] * (n - len(row))


def _url_formulas(ws, cols: dict[str, int]) -> list[list]:
    letter = get_column_letter(cols["URL"])
    return ws.get(f"{letter}:{letter}", value_render_option="FORMULA")  # =HYPERLINK("target", "text")


# ---------- duplicate check ----------

def _job_rows(values: list[list], formulas: list[list], cols: dict[str, int]) -> list[tuple[int, dict, list[str]]]:
    """[(row number, {field: display value}, urls in the URL cell)] like tracker.read_rows; rows with no
    Company and no URL are skipped."""
    width = max(cols.values())
    rows = []
    for i, raw in enumerate(values[1:], start=2):
        padded = _pad(raw, width)
        row = {f: padded[cols[f] - 1] for f in tracker.FIELDS}
        if not row["Company"] and not row["URL"]:
            continue
        urls = [row["URL"]] if row["URL"] else []
        formula = _pad(formulas[i - 1], 1)[0] if i - 1 < len(formulas) else ""
        m = HYPERLINK_TARGET.match(str(formula))
        if m:
            urls.append(m.group(1).replace('""', '"'))
        rows.append((i, row, urls))
    return rows


@_guard
def read_rows(ws) -> list[tuple[int, dict, list[str]]]:
    """Every job row of the Sheet: [(row number, {field: display value}, urls in the URL cell)]."""
    values = ws.get_all_values()
    cols = columns(values)
    return _job_rows(values, _url_formulas(ws, cols), cols)


@_guard
def find_duplicates(ws, url: str | None = None, company: str | None = None,
                    position: str | None = None) -> list[tuple[int, dict]]:
    """Same matching as tracker.find_duplicates, on the Sheet. -> [(row, display values)]"""
    return tracker.match_rows(read_rows(ws), url, company, position)


# ---------- writing ----------

def next_row(values: list[list], company_col: int = 1) -> int:
    """First row after the last row that has a Company value."""
    for i in range(len(values) - 1, 0, -1):
        if str(_pad(values[i], company_col)[company_col - 1]).strip():
            return i + 2
    return 2


def backup_values(values: list[list], backup_dir) -> Path:
    """Save the Sheet's values to backups/sheet-YYYYMMDD-HHMMSS.csv before any write."""
    backup_dir = Path(backup_dir)
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = backup_dir / f"sheet-{stamp}.csv"
    n = 1
    while target.exists():  # two writes in the same second
        target = backup_dir / f"sheet-{stamp}-{n}.csv"
        n += 1
    with open(target, "w", newline="", encoding="utf-8-sig") as f:  # -sig: Excel reads it as UTF-8
        csv.writer(f).writerows(values)
    return target


def _quote(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def row_values(row: TrackerRow) -> dict[str, str]:
    """The new row's cells by field. The ISO dates are understood in every locale."""
    return {
        "Company": row.company,
        "Position": row.position,
        "URL": f"=HYPERLINK({_quote(row.url)}, {_quote(row.url)})",
        "Posted date/ Source": row.posted,
        "Date": row.applied_on.isoformat(),
        "Outcome": tracker.OUTCOME,
        "Last Updated": row.applied_on.isoformat(),
    }


def _write_cells(ws, r: int, cells: dict[str, str], cols: dict[str, int]) -> None:
    """Write each field into its own cell of row r: text RAW, dates and the URL formula USER_ENTERED."""
    for option, fields in (("RAW", [f for f in cells if f not in ENTERED]),
                           ("USER_ENTERED", [f for f in cells if f in ENTERED])):
        if fields:
            ws.batch_update([{"range": f"{get_column_letter(cols[f])}{r}", "values": [[cells[f]]]} for f in fields],
                            value_input_option=option)


def _date_format(sheet_id: int, r: int, col: int) -> dict:
    """Show the cell at row r, column col (both 1-based; the API's are 0-based) as M/d/yyyy."""
    return {"repeatCell": {
        "range": {"sheetId": sheet_id, "startRowIndex": r - 1, "endRowIndex": r,
                  "startColumnIndex": col - 1, "endColumnIndex": col},
        "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": DATE_PATTERN}}},
        "fields": "userEnteredFormat.numberFormat",
    }}


def _format_requests(sheet_id: int, r: int, cols: dict[str, int]) -> list[dict]:
    """Copy format + data validation (Outcome and Stage dropdowns, date checks) from row r-1 to row r,
    then give the date columns the M/d/yyyy date format. Rows here are 1-based; the API's are 0-based."""
    def grid(row):
        return {"sheetId": sheet_id, "startRowIndex": row - 1, "endRowIndex": row,
                "startColumnIndex": 0, "endColumnIndex": max(cols.values())}
    requests = []
    if r > 2:
        for paste in ("PASTE_FORMAT", "PASTE_DATA_VALIDATION"):
            requests.append({"copyPaste": {"source": grid(r - 1), "destination": grid(r), "pasteType": paste}})
    requests += [_date_format(sheet_id, r, cols[f]) for f in tracker.DATE_FIELDS]
    return requests


@_guard
def append_row(ws, row: TrackerRow, backup_dir) -> int:
    """Check the header, back up, then write `row` after the last Company row. Returns the row number."""
    values = ws.get_all_values()
    cols = columns(values)
    r = next_row(values, cols["Company"])
    backup_values(values, backup_dir)
    if r > ws.row_count:
        ws.add_rows(SHEET_GROWTH)
    _write_cells(ws, r, row_values(row), cols)
    ws.spreadsheet.batch_update({"requests": _format_requests(ws.id, r, cols)})
    return r


# ---------- Outcome updates (find.py) ----------

def ensure_choices(ws, col: int, name: str, wanted) -> None:
    """Make the Sheet's dropdown on column `col` (its row-2 list rule) offer every one of `wanted`.
    Only a plain list rule is changed, and only by adding; any other rule gets a note and is left alone."""
    letter = get_column_letter(col)
    meta = ws.spreadsheet.fetch_sheet_metadata(params={
        "ranges": f"'{ws.title}'!{letter}2", "includeGridData": "true",
        "fields": "sheets(properties.sheetId,data.rowData.values.dataValidation)"})
    try:
        cell = meta["sheets"][0]["data"][0]["rowData"][0]["values"][0]
    except (KeyError, IndexError):
        return
    rule = cell.get("dataValidation")
    if not rule:
        return
    condition = rule.get("condition", {})
    if condition.get("type") != "ONE_OF_LIST":
        print(f"Note: the Sheet's {name} dropdown is a {condition.get('type')} rule; "
              f"check it offers {', '.join(wanted)}.")
        return
    choices = [v.get("userEnteredValue", "") for v in condition.get("values", [])]
    missing = [o for o in wanted if o not in choices]
    if not missing:
        return
    new_rule = {**rule, "condition": {**condition, "values": [{"userEnteredValue": c} for c in choices + missing]}}
    ws.spreadsheet.batch_update({"requests": [{"setDataValidation": {
        "range": {"sheetId": ws.id, "startRowIndex": 1, "endRowIndex": ws.row_count,
                  "startColumnIndex": col - 1, "endColumnIndex": col},
        "rule": new_rule,
    }}]})


def ensure_dropdowns(ws, cols: dict[str, int]) -> None:
    ensure_choices(ws, cols["Outcome"], "Outcome", tracker.OUTCOMES)
    ensure_choices(ws, cols["Stage"], "Stage", tracker.STAGES)


@_guard
def update_outcome(ws, updates: list[tracker.OutcomeUpdate], backup_dir) -> tuple[list, list]:
    """Same as tracker.update_outcome, on the Sheet: check the header, back up, then write Outcome
    (and Stage / Remarks when given) row by row. -> ([(row, update) written], [not found])"""
    values = ws.get_all_values()
    cols = columns(values)
    rows = _job_rows(values, _url_formulas(ws, cols), cols)
    found = [(tracker.locate(rows, u), u) for u in updates]
    written = [(r, u) for r, u in found if r is not None]
    missing = [u for r, u in found if r is None]
    if not written:
        return written, missing
    backup_values(values, backup_dir)
    current = {r: v for r, v, _ in rows}  # kept up to date, so a second update to a row builds on the first
    for r, u in written:
        cells = {"Outcome": u.outcome}
        if u.stage:
            cells["Stage"] = u.stage
        if u.note.strip():
            cells["Remarks"] = current[r]["Remarks"] = tracker.format_note(u.noted_on, u.note, current[r]["Remarks"])
        _write_cells(ws, r, cells, cols)
    ensure_dropdowns(ws, cols)
    return written, missing
